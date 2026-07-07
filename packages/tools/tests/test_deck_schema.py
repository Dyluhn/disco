"""Tests for C1 — _deck_schema.py.

Proves:
  - lower_deck produces a valid Deck from a sample AuthoredDeck.
  - _fit_text steps font down when text overflows.
  - Overflow body creates a continuation slide (NOT truncation).
  - _infer_layout is total (never returns None, handles unknown types).
  - Theme resolution splits "disco-light" → ("disco", "light") correctly.
  - image_prompt + empty body → full_image layout.
  - image_prompt + body → image_right layout.
  - notes are NOT counted as overflow (excluded from _fit_text).
  - All layout functions produce in-canvas elements (left/top/width/height > 0).
  - image_prompt element carries the prompt for C7 wiring.
"""

from __future__ import annotations

from disco.tools.builtin._deck_schema import (
    _CH,
    _CW,
    _SLIDE_H,
    _SLIDE_W,
    AuthoredDeck,
    AuthoredSlide,
    ChartSpec,
    Deck,
    TableSpec,
    _fit_text,
    _infer_layout,
    _parse_theme,
    lower_deck,
    lower_deck_for_editor,
)

# ---------------------------------------------------------------------------
# Sample fixtures
# ---------------------------------------------------------------------------

def _sample_authored_deck() -> AuthoredDeck:
    """A representative 5-slide AuthoredDeck covering multiple layouts."""
    return AuthoredDeck(
        title="Disco AI — Product Launch",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="title",
                title="Introducing Disco AI",
                body=["The research assistant that cites its sources"],
                notes="Welcome the audience warmly.",
            ),
            AuthoredSlide(
                type="bullets",
                title="Key Features",
                body=[
                    "Grounded answers with citation accuracy > 90%",
                    "Multi-source deep research in under 60 seconds",
                    "Privacy-first: runs fully on your hardware",
                    "Open-source core with permissive licensing",
                ],
            ),
            AuthoredSlide(
                type="section_header",
                title="How It Works",
                body=["Three-stage pipeline: search → extract → ground"],
            ),
            AuthoredSlide(
                type="image_right",
                title="Architecture Overview",
                body=["Query router", "Parallel retrieval", "NLI grounding"],
                image_prompt="A clean technical architecture diagram with three boxes connected by arrows on a light background",
            ),
            AuthoredSlide(
                type="closing",
                title="Thank You",
                body=["disco.ai | hello@disco.ai"],
            ),
        ],
    )


def _overflow_authored_deck() -> AuthoredDeck:
    """A deck with a slide that has 30 bullet lines — must trigger continuation."""
    return AuthoredDeck(
        title="Overflow Test",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="title",
                title="Overflow Demo",
                body=["Testing continuation slides"],
            ),
            AuthoredSlide(
                type="bullets",
                title="Dense Slide",
                body=[f"Bullet point number {i+1}: some reasonably detailed content here" for i in range(30)],
            ),
        ],
    )


def _image_deck() -> AuthoredDeck:
    """Deck with image slides to test full_image vs image_right."""
    return AuthoredDeck(
        title="Image Test",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="full_image",
                title="Urban Mobility 2035",
                body=[],  # no body → full_image
                image_prompt="Futuristic cityscape with autonomous electric vehicles at sunset",
            ),
            AuthoredSlide(
                type="image_right",
                title="Design Philosophy",
                body=["Human-centred", "Sustainable materials", "Zero emissions"],
                image_prompt="Minimal product design sketch on white background",
            ),
        ],
    )


# ---------------------------------------------------------------------------
# lower_deck — basic structure
# ---------------------------------------------------------------------------


def test_lower_deck_returns_deck():
    """lower_deck returns a Deck instance."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    assert isinstance(deck, Deck)


def test_lower_deck_title_preserved():
    """The deck title is preserved from the AuthoredDeck."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    assert deck.title == authored.title


def test_lower_deck_theme_resolved():
    """Theme tokens are resolved — disco-light → non-None bg/text/accent."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    assert deck.theme is not None
    assert deck.theme.bg.startswith("#")
    assert deck.theme.accent.startswith("#")
    assert deck.theme.text.startswith("#")


def test_lower_deck_produces_slides():
    """At least as many slides as the authored deck (continuation may add more)."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    assert len(deck.slides) >= len(authored.slides)


def test_lower_deck_slides_have_layouts():
    """Every lowered slide has a non-empty, non-None layout."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    for slide in deck.slides:
        assert slide.layout
        assert isinstance(slide.layout, str)


def test_lower_deck_slides_have_elements():
    """Every lowered slide has at least one Element."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    for slide in deck.slides:
        assert len(slide.elements) >= 1


def test_lower_deck_notes_on_slide():
    """Notes from AuthoredSlide appear on the lowered Slide (not as an element)."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    # First slide has notes
    first = deck.slides[0]
    assert first.notes == "Welcome the audience warmly."
    # No element should contain the notes text
    all_text = " ".join(el.text for el in first.elements if el.kind == "text")
    assert "Welcome the audience warmly." not in all_text


def test_lower_deck_16_9_canvas():
    """Deck size is 12_192_000 × 6_858_000 EMU (16:9)."""
    deck = lower_deck(_sample_authored_deck())
    assert deck.size == (12_192_000, 6_858_000)


# ---------------------------------------------------------------------------
# Overflow → continuation split (NOT truncation)
# ---------------------------------------------------------------------------


def test_overflow_creates_continuation_slide():
    """A 30-bullet slide MUST produce at least 2 slides (original + continuation)."""
    authored = _overflow_authored_deck()
    deck = lower_deck(authored)

    # The overflow deck has 2 authored slides; the dense one should split.
    dense_slides = [s for s in deck.slides if "Dense" in s.type or "_cont" in s.type or s.type == "bullets"]
    # At minimum, the original "Dense Slide" and one continuation
    assert len(deck.slides) >= 3, (
        f"Expected at least 3 slides (title + dense + continuation), got {len(deck.slides)}"
    )


def test_overflow_no_truncation():
    """All 30 bullet texts must appear across all slides — none silently dropped."""
    authored = _overflow_authored_deck()
    deck = lower_deck(authored)

    # Collect all text from elements across all slides
    all_text = " ".join(
        el.text for slide in deck.slides
        for el in slide.elements
        if el.kind == "text"
    )
    for i in range(30):
        expected = f"Bullet point number {i+1}"
        assert expected in all_text, (
            f"Bullet {i+1} was truncated (not found in any slide element text)"
        )


def test_continuation_slide_has_cont_type():
    """Continuation slides have type ending in '_cont'."""
    authored = _overflow_authored_deck()
    deck = lower_deck(authored)
    cont_slides = [s for s in deck.slides if "_cont" in s.type]
    assert len(cont_slides) >= 1, "No continuation slide found after overflow"


def test_photo_grid_two_line_slide_does_not_create_orphan_continuation():
    """A two-line image/photo-grid slide must not split one line into a bare cont slide."""
    authored = AuthoredDeck(
        title="Photo Grid Orphan",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="photo_grid",
                archetype="photo_grid",
                title="Field Evidence",
                body=["Line one caption", "Line two caption"],
                layout_hint="full_image",
                image_prompt="documentary grid of field evidence",
            )
        ],
    )

    deck = lower_deck(authored)
    editor = lower_deck_for_editor(authored)
    all_text = " ".join(
        el.text for slide in deck.slides for el in slide.elements if el.kind == "text"
    )

    assert len(deck.slides) == 1
    assert len(editor.slides) == 1
    assert not any(slide.title.endswith("(cont.)") for slide in deck.slides)
    assert "Line one caption" in all_text
    assert "Line two caption" in all_text


# ---------------------------------------------------------------------------
# BW-12 / BW-13 — editor and export agree; no duplicate title pages
# ---------------------------------------------------------------------------


def _title_overflow_deck() -> AuthoredDeck:
    """A title slide carrying MORE than one body line — the BW-12 duplicate trap.

    The old lowerer spilled body[1:] into a ``title_cont`` that re-inferred back
    to the ``title`` layout → a SECOND identical title page.
    """
    return AuthoredDeck(
        title="Dup Test",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="title",
                title="Welcome",
                body=["Subtitle line", "Spilled line A", "Spilled line B"],
            ),
        ],
    )


def test_continuation_never_renders_a_second_title_layout():
    """BW-12: a title slide that overflows must NOT produce two 'title' slides."""
    deck = lower_deck(_title_overflow_deck())
    title_layouts = [s for s in deck.slides if s.layout == "title"]
    assert len(title_layouts) == 1, (
        f"Expected exactly one title-layout slide, got {len(title_layouts)} "
        f"(layouts={[s.layout for s in deck.slides]})"
    )
    # The overflow continuation is a bullets slide relabelled '(cont.)'
    conts = [s for s in deck.slides if "_cont" in s.type]
    assert conts and all(s.layout == "bullets" for s in conts)
    assert all(s.title.endswith("(cont.)") for s in conts)


def test_editor_count_matches_export_count():
    """BW-13: lower_deck_for_editor must produce the SAME slide count as lower_deck."""
    for authored in (_overflow_authored_deck(), _title_overflow_deck(), _sample_authored_deck()):
        exported = lower_deck(authored)
        editor = lower_deck_for_editor(authored)
        assert len(editor.slides) == len(exported.slides), (
            f"editor={len(editor.slides)} vs export={len(exported.slides)} "
            f"for deck {authored.title!r}"
        )


def test_editor_continuation_bullets_point_at_original_body():
    """BW-13: spilled bullets in the editor map back to real /slides/i/body/j pointers."""
    editor = lower_deck_for_editor(_overflow_authored_deck())
    # Gather every authored body index referenced across all (split) editor slides
    referenced: set[str] = set()
    for s in editor.slides:
        for el in s.elements:
            if el.kind in ("bullet", "subtitle"):
                referenced.add(el.json_pointer)
    # The dense slide is authored index 1 with 30 body lines — all must be addressable
    for j in range(30):
        assert f"/slides/1/body/{j}" in referenced, f"body/{j} not addressable in editor"


# ---------------------------------------------------------------------------
# BW-13 — NON-SUFFIX overflow (column layouts spill non-contiguously)
# ---------------------------------------------------------------------------

import re  # noqa: E402


def _two_column_overflow_deck() -> AuthoredDeck:
    """A two_column slide dense enough to overflow at the font floor.

    Unlike bullets (which spills a contiguous TAIL), two_column drops the TAIL of
    EACH column, so the overflow is a non-contiguous union — the case where the
    old ``body[:consumed]`` editor model mis-mapped pointers.
    """
    body = [f"Line {i + 1}: " + ("content " * 14) for i in range(40)]
    return AuthoredDeck(
        title="Two-Column Overflow",
        theme="disco-light",
        slides=[AuthoredSlide(type="two_column", title="Dense Columns", body=body)],
    )


def _comparison_overflow_deck() -> AuthoredDeck:
    """A comparison slide (body[0]/body[1] labels + dense 50/50 content)."""
    body = ["LEFT SIDE", "RIGHT SIDE"] + [
        f"Item {i + 1}: " + ("detail " * 14) for i in range(40)
    ]
    return AuthoredDeck(
        title="Comparison Overflow",
        theme="disco-light",
        slides=[AuthoredSlide(type="comparison", title="Dense Compare", body=body)],
    )


def _assert_editor_pointer_parity(authored: AuthoredDeck, orig: int, nbody: int) -> None:
    """Editor count == export count AND every editor bullet/subtitle points at the
    AUTHORED body line whose text it displays — for ALL overflow kinds (BW-13)."""
    exported = lower_deck(authored)
    editor = lower_deck_for_editor(authored)

    # Count parity (must also have actually split — multi-slide).
    assert len(editor.slides) == len(exported.slides), (
        f"editor={len(editor.slides)} vs export={len(exported.slides)}"
    )
    assert len(editor.slides) >= 2, "fixture did not overflow — not exercising the split"

    body = authored.slides[orig].body
    seen: dict[int, int] = {}
    for s in editor.slides:
        for el in s.elements:
            if el.kind not in ("bullet", "subtitle"):
                continue
            m = re.search(rf"/slides/{orig}/body/(\d+)$", el.json_pointer)
            assert m is not None, f"unexpected pointer {el.json_pointer}"
            j = int(m.group(1))
            # The pointer must reference the authored line this element DISPLAYS.
            assert el.content == body[j].lstrip("• "), (
                f"pointer {el.json_pointer} shows {el.content!r} "
                f"but body[{j}]={body[j]!r}"
            )
            seen[j] = seen.get(j, 0) + 1

    # Every authored body line addressable exactly once across the split.
    assert sorted(seen) == list(range(nbody)), (
        f"addressed {sorted(seen)} != 0..{nbody - 1}"
    )
    assert all(v == 1 for v in seen.values()), f"duplicate pointers: {seen}"


def test_editor_two_column_non_suffix_overflow_pointer_parity():
    """BW-13: two_column overflow (non-contiguous spill) — editor pointers stay correct."""
    _assert_editor_pointer_parity(_two_column_overflow_deck(), orig=0, nbody=40)


def test_editor_comparison_non_suffix_overflow_pointer_parity():
    """BW-13: comparison overflow (labels + per-column tails) — editor pointers stay correct."""
    _assert_editor_pointer_parity(_comparison_overflow_deck(), orig=0, nbody=42)


def _bullets_overflow_deck() -> AuthoredDeck:
    """A plain bullets slide dense enough to spill a contiguous TAIL into a continuation."""
    body = [f"Bullet {i + 1}: " + ("word " * 16) for i in range(40)]
    return AuthoredDeck(
        title="Bullets Overflow",
        theme="disco-light",
        slides=[AuthoredSlide(type="bullets", title="Dense Bullets", body=body)],
    )


def _render_body_ids_by_slide(authored: AuthoredDeck) -> dict[str, dict[str, str]]:
    """Parse render_html → {slide_id: {data-element-id: displayed_text}} for body lines.

    This is what the REAL editor consumes: SlideCanvas joins each rendered
    ``[data-element-id]`` node to the editor pointer model by that exact id. We pull
    the ``{sid}:body:{j}`` ``<li>`` elements (id + their visible text) so a test can
    assert the rendered identity equals the pointer model, not merely the count.
    """
    from disco.tools.builtin._pptx_render import render_html

    html_str = render_html(lower_deck(authored))
    out: dict[str, dict[str, str]] = {}
    # <li data-element-id="slide-N:body:J" data-slide-id="slide-N">TEXT</li>
    for eid, sid, text in re.findall(
        r'<li data-element-id="([^"]+)" data-slide-id="([^"]+)">(.*?)</li>',
        html_str,
        re.DOTALL,
    ):
        out.setdefault(sid, {})[eid] = text
    return out


def _assert_render_pointer_identity_alignment(authored: AuthoredDeck, orig: int) -> None:
    """BW-13 residual P1: the RENDERED editor slides (render_html, what DeckExportBar /
    SlideCanvas key on) and the editor POINTER model (lower_deck_for_editor) must share
    the SAME slide identity AND the SAME per-body-line element ids — so each rendered
    editor slide maps to the correct authored line in the actual UI, not just in counts.

    Asserted for overflow/continuation slides (where render position != authored index):
      • identical set of slide ids,
      • for every rendered ``{sid}:body:{orig_j}`` <li>, the editor model has the SAME
        element id with a json_pointer to ``/slides/{orig}/body/{orig_j}``, AND
      • the authored line that pointer addresses is the line the rendered <li> DISPLAYS.
    """
    editor = lower_deck_for_editor(authored)
    rendered = _render_body_ids_by_slide(authored)
    body = authored.slides[orig].body

    # Must actually have split — otherwise we are not exercising continuation identity.
    assert len(editor.slides) >= 2, "fixture did not overflow — not exercising the split"

    # Same slide identity/order on both sides.
    assert [s.slide_id for s in editor.slides] == sorted(
        rendered, key=lambda s: int(s.split("-")[1])
    ), "rendered slide ids != editor slide ids"

    # Editor body element_id → json_pointer, keyed by slide id.
    editor_body: dict[str, dict[str, str]] = {}
    for s in editor.slides:
        editor_body[s.slide_id] = {
            el.element_id: el.json_pointer for el in s.elements if el.kind == "bullet"
        }

    for sid, rendered_lis in rendered.items():
        for eid, displayed in rendered_lis.items():
            # The rendered <li>'s id must exist VERBATIM in the editor pointer model.
            assert eid in editor_body[sid], (
                f"rendered {eid} on {sid} absent from editor model {sorted(editor_body[sid])}"
            )
            ptr = editor_body[sid][eid]
            m = re.search(rf"/slides/{orig}/body/(\d+)$", ptr)
            assert m is not None, f"unexpected pointer {ptr} for {eid}"
            j = int(m.group(1))
            # The id encodes the authored index; the pointer must agree with it...
            assert eid.endswith(f":body:{j}"), f"id {eid} disagrees with pointer {ptr}"
            # ...and the authored line that index addresses is what the slide DISPLAYS.
            assert displayed == body[j].lstrip("• "), (
                f"{eid} shows {displayed!r} but /body/{j}={body[j]!r}"
            )


def test_render_pointer_identity_two_column_overflow():
    """BW-13 P1: two_column non-suffix overflow — rendered ids == editor pointer ids."""
    _assert_render_pointer_identity_alignment(_two_column_overflow_deck(), orig=0)


def test_render_pointer_identity_comparison_overflow():
    """BW-13 P1: comparison non-suffix overflow — rendered ids == editor pointer ids."""
    _assert_render_pointer_identity_alignment(_comparison_overflow_deck(), orig=0)


def test_render_pointer_identity_bullets_continuation_overflow():
    """BW-13 P1: bullets SUFFIX overflow — continuation-slide rendered ids == editor ids.

    The continuation slide's bullets render at LOCAL positions 0,1,2… but address the
    authored TAIL (body[k], body[k+1]…). Before the fix render_html stamped the local
    position, so the rendered <li> id (``body:0``) never matched the editor pointer
    (``body:21``) — the overlay could not bind. This pins them equal.
    """
    _assert_render_pointer_identity_alignment(_bullets_overflow_deck(), orig=0)


# ---------------------------------------------------------------------------
# _fit_text — font step-down
# ---------------------------------------------------------------------------


def test_fit_text_short_content_no_overflow():
    """4 lines of short bullets fit at max font."""
    from disco.tools.builtin._deck_schema import _BODY_FONT_MAX, _CH, _CW
    lines = ["Short bullet one", "Short bullet two", "Short bullet three", "Short bullet four"]
    font, fitted, overflow = _fit_text(lines, box_width=float(_CW), box_height=float(_CH))
    assert overflow == []
    assert fitted == lines
    assert font == _BODY_FONT_MAX


def test_fit_text_many_long_lines_steps_down():
    """Many long lines force font step-down below max."""
    from disco.tools.builtin._deck_schema import _BODY_FONT_MAX
    lines = ["A" * 150 for _ in range(20)]  # very long lines, many of them
    font, fitted, overflow = _fit_text(lines, box_width=float(_CW), box_height=float(_CH))
    # Should have stepped down OR produced overflow
    assert font < _BODY_FONT_MAX or len(overflow) > 0


def test_fit_text_overflow_is_not_truncated():
    """When overflow occurs, fitted + overflow == all input lines."""
    lines = [f"Bullet line {i}" for i in range(50)]
    font, fitted, overflow = _fit_text(lines, box_width=float(_CW), box_height=float(_CH))
    assert fitted + overflow == lines


def test_fit_text_empty_input():
    """Empty input returns (max_font, [], [])."""
    from disco.tools.builtin._deck_schema import _BODY_FONT_MAX
    font, fitted, overflow = _fit_text([], box_width=float(_CW), box_height=float(_CH))
    assert font == _BODY_FONT_MAX
    assert fitted == []
    assert overflow == []


# ---------------------------------------------------------------------------
# _infer_layout — total + correct
# ---------------------------------------------------------------------------


def test_infer_layout_image_prompt_no_body_is_full_image():
    """image_prompt + no body → full_image."""
    slide = AuthoredSlide(type="image_right", title="X", body=[], image_prompt="a mountain")
    assert _infer_layout(slide) == "full_image"


def test_infer_layout_image_prompt_with_body_is_image_right():
    """image_prompt + body → image_right (first call, alt counter=0)."""
    slide = AuthoredSlide(type="bullets", title="X", body=["bullet"], image_prompt="a forest")
    counter = [0]
    layout = _infer_layout(slide, counter)
    assert layout in ("image_right", "image_left")


def test_infer_layout_alternates_image_sides():
    """Alternating image_prompt+body slides get image_right then image_left."""
    counter = [0]
    s = AuthoredSlide(type="bullets", title="X", body=["b"], image_prompt="img")
    l1 = _infer_layout(s, counter)
    l2 = _infer_layout(s, counter)
    assert l1 != l2
    assert set([l1, l2]) == {"image_right", "image_left"}


def test_infer_layout_title():
    assert _infer_layout(AuthoredSlide(type="title", title="T")) == "title"


def test_infer_layout_section_header():
    assert _infer_layout(AuthoredSlide(type="section_header", title="S")) == "section_header"
    assert _infer_layout(AuthoredSlide(type="section", title="S")) == "section_header"


def test_infer_layout_closing():
    assert _infer_layout(AuthoredSlide(type="closing", title="C")) == "closing"


def test_infer_layout_two_column():
    assert _infer_layout(AuthoredSlide(type="two_column", title="T")) == "two_column"


def test_infer_layout_unknown_falls_back_to_bullets():
    """Unknown custom type → bullets (never None, never crash)."""
    slide = AuthoredSlide(type="era_1_timeline", title="T", body=["b"])
    layout = _infer_layout(slide)
    assert layout == "bullets"


def test_infer_layout_cont_type_is_bullets():
    """Continuation type like 'bullets_cont' → bullets."""
    slide = AuthoredSlide(type="bullets_cont", title="T", body=["b"])
    assert _infer_layout(slide) == "bullets"


def test_infer_layout_layout_hint_wins():
    """Explicit layout_hint always overrides type-based inference."""
    slide = AuthoredSlide(type="title", title="T", layout_hint="closing")
    assert _infer_layout(slide) == "closing"


def test_infer_layout_chart_triggers_metrics():
    """chart set → metrics layout."""
    slide = AuthoredSlide(
        type="bullets",
        title="Chart",
        chart=ChartSpec(kind="bar", labels=["A", "B"], series=[{"name": "x", "data": [1.0, 2.0]}]),
    )
    assert _infer_layout(slide) == "metrics"


def test_infer_layout_table_triggers_table():
    """table set → table layout."""
    slide = AuthoredSlide(
        type="bullets",
        title="Table",
        table=TableSpec(headers=["A", "B"], rows=[["1", "2"]]),
    )
    assert _infer_layout(slide) == "table"


# ---------------------------------------------------------------------------
# Theme parsing
# ---------------------------------------------------------------------------


def test_parse_theme_disco_light():
    assert _parse_theme("disco-light") == ("disco", "light")


def test_parse_theme_disco_dark():
    assert _parse_theme("disco-dark") == ("disco", "dark")


def test_parse_theme_neutral():
    assert _parse_theme("neutral") == ("neutral", "light")


# ---------------------------------------------------------------------------
# Image wiring in lowered elements
# ---------------------------------------------------------------------------


def test_image_element_has_prompt():
    """An image slide produces an Element(kind='image') with the image_prompt."""
    authored = _image_deck()
    deck = lower_deck(authored)
    img_elements = [
        el
        for slide in deck.slides
        for el in slide.elements
        if el.kind == "image"
    ]
    assert len(img_elements) >= 1
    assert any(el.image_prompt for el in img_elements), (
        "No image element carries the image_prompt for C7 wiring"
    )


def test_full_image_layout_for_no_body():
    """image_prompt + body=[] → full_image layout on the lowered slide."""
    authored = _image_deck()
    deck = lower_deck(authored)
    full_img_slides = [s for s in deck.slides if s.layout == "full_image"]
    assert len(full_img_slides) >= 1


def test_notes_excluded_from_elements():
    """Notes are not rendered as text elements on the slide face."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    for slide in deck.slides:
        for el in slide.elements:
            if el.kind == "text" and slide.notes:
                assert slide.notes not in el.text, (
                    f"Notes text leaked into a visible element on slide {slide.type}"
                )


# ---------------------------------------------------------------------------
# In-canvas element geometry
# ---------------------------------------------------------------------------


def test_all_elements_in_canvas():
    """Every element's bounding box is within the 16:9 canvas."""
    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    for slide in deck.slides:
        for el in slide.elements:
            assert el.left >= 0, f"Element left={el.left} < 0"
            assert el.top >= 0, f"Element top={el.top} < 0"
            assert el.left + el.width <= _SLIDE_W + 1000, (
                f"Element right={el.left + el.width} exceeds canvas {_SLIDE_W}"
            )
            assert el.top + el.height <= _SLIDE_H + 1000, (
                f"Element bottom={el.top + el.height} exceeds canvas {_SLIDE_H}"
            )


# ---------------------------------------------------------------------------
# Deck round-trip through PPTX renderer
# ---------------------------------------------------------------------------


def test_lower_deck_renders_to_valid_pptx():
    """A lowered deck can be rendered to a valid .pptx file by C3."""
    import io

    from disco.tools.builtin._pptx_render import render_pptx
    from pptx import Presentation

    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    pptx_bytes = render_pptx(deck)
    assert pptx_bytes[:4] == b"PK\x03\x04", "Not a valid ZIP/OOXML"
    prs = Presentation(io.BytesIO(pptx_bytes))
    # Slide count may be > authored count due to continuations
    assert len(prs.slides) == len(deck.slides)


def test_lower_deck_renders_title_text_in_pptx():
    """Slide titles appear in the .pptx text frames."""
    import io

    from disco.tools.builtin._pptx_render import render_pptx
    from pptx import Presentation

    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    pptx_bytes = render_pptx(deck)
    prs = Presentation(io.BytesIO(pptx_bytes))

    all_pptx_text = ""
    for pptx_slide in prs.slides:
        for shape in pptx_slide.shapes:
            if shape.has_text_frame:
                all_pptx_text += shape.text_frame.text + " "

    assert "Introducing Disco AI" in all_pptx_text
    assert "Key Features" in all_pptx_text


def test_lower_deck_renders_to_valid_html():
    """A lowered deck renders to valid 16:9 HTML."""
    from disco.tools.builtin._pptx_render import render_html

    authored = _sample_authored_deck()
    deck = lower_deck(authored)
    html_str = render_html(deck)

    assert "<!DOCTYPE html>" in html_str
    assert "disco-slides-generate:pipeline-html" in html_str
    assert "56.25vw" in html_str  # 16:9 ratio
    assert "ArrowRight" in html_str  # keyboard nav
    assert html_str.count('<section class="slide') == len(deck.slides)


def test_archetype_html_has_distinctive_layout_structures():
    from disco.tools.builtin._pptx_render import render_html

    authored = AuthoredDeck(
        title="Archetype Render",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="bullets",
                archetype="big_number",
                title="Adoption",
                body=["73%", "of teams ship the pilot within one week"],
            ),
            AuthoredSlide(
                type="bullets",
                archetype="quote",
                title="Customer Voice",
                body=["The new flow removed the review bottleneck.", "Avery Lee, Ops"],
            ),
            AuthoredSlide(
                type="bullets",
                archetype="timeline",
                title="Rollout",
                body=["Q1: Pilot", "Q2: Integrate", "Q3: Scale"],
            ),
            AuthoredSlide(
                type="bullets",
                archetype="two_by_two",
                title="Decision Map",
                body=["Fast wins", "Strategic bets", "Maintenance", "Avoid"],
            ),
        ],
    )

    html_str = render_html(lower_deck(authored))

    assert 'data-archetype="big_number"' in html_str
    assert 'class="slide-big-number-figure"' in html_str
    assert 'class="slide-quote-mark"' in html_str
    assert html_str.count('class="slide-timeline-node"') == 3
    assert html_str.count('class="slide-two-by-two-cell"') == 4


def test_two_by_two_axis_labels_only_when_authored() -> None:
    """No-false-affordances (gauntlet e-news 2026-07-07): a 4-cell two_by_two must
    NOT get invented 'Higher impact/certainty' axis labels — the renderer was
    fabricating an analytical framing the quadrants never plotted. Axis labels
    render ONLY when the model authored them (6 body lines: x, y, 4 cells)."""
    from disco.tools.builtin._deck_schema import AuthoredDeck, AuthoredSlide, lower_deck
    from disco.tools.builtin._pptx_render import render_html, render_pptx

    def _deck(body: list[str]) -> AuthoredDeck:
        return AuthoredDeck(
            title="T", theme="disco-light",
            slides=[AuthoredSlide(type="bullets", archetype="two_by_two",
                                  title="Decision Map", body=body)],
        )

    four = _deck(["Fast wins", "Strategic bets", "Maintenance", "Avoid"])
    html_4 = render_html(lower_deck(four))
    assert "Higher impact" not in html_4 and "Higher certainty" not in html_4
    assert html_4.count('class="slide-two-by-two-cell"') == 4

    six = _deck(["Effort", "Value", "Fast wins", "Strategic bets", "Maintenance", "Avoid"])
    html_6 = render_html(lower_deck(six))
    assert "Effort" in html_6 and "Value" in html_6  # authored axes DO render

    # PPTX path: no invented axis text on the 4-cell slide.
    import io

    from pptx import Presentation
    prs = Presentation(io.BytesIO(render_pptx(lower_deck(four))))
    all_text = " ".join(
        s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame
    )
    assert "Higher impact" not in all_text and "Higher certainty" not in all_text


def test_lower_deck_theme_override_rethemes_without_mutating_authored() -> None:
    """The slide-deck template selector path: theme_override re-themes at render
    time, the authored deck's own theme is untouched, and an unknown id raises."""
    import pytest

    from disco.core.brand import resolve_theme
    from disco.tools.builtin._deck_schema import AuthoredDeck, AuthoredSlide, lower_deck

    authored = AuthoredDeck(
        title="T", theme="disco-light",
        slides=[AuthoredSlide(type="title", title="Hi", body=["x"])],
    )
    base = lower_deck(authored)
    over = lower_deck(authored, theme_override="midnight-dark")
    assert base.theme == resolve_theme("disco", "light")
    assert over.theme == resolve_theme("midnight", "dark")
    assert authored.theme == "disco-light"  # authored sidecar NOT mutated
    with pytest.raises(ValueError):
        lower_deck(authored, theme_override="bogus-template")


def test_lower_deck_brand_override_uses_direction_tokens() -> None:
    from disco.core.design import DIRECTION_BY_ID, to_brand_tokens
    from disco.tools.builtin._deck_schema import AuthoredDeck, AuthoredSlide, lower_deck

    authored = AuthoredDeck(
        title="T", theme="disco-light",
        slides=[AuthoredSlide(type="title", title="Hi", body=["x"])],
    )
    theme = to_brand_tokens(DIRECTION_BY_ID["brutalist"])
    deck = lower_deck(authored, brand_override=theme)

    assert deck.theme == theme
    assert deck.theme.accent == "#c23616"
    assert authored.theme == "disco-light"
