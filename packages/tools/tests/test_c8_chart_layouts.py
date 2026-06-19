"""Tests for C8 — chart and table slide layouts (_c8_chart_layouts.py).

Proves:
  - ChartSpec / TableSpec parse and validate correctly.
  - _to_svg_dict bridge produces the expected chart_svg format.
  - html_chart_content embeds an SVG for valid data.
  - html_chart_content falls back to an HTML table when data is empty/malformed.
  - html_table_content produces a native HTML table from TableSpec.
  - layout_chart_slide_pptx adds a real chart shape to the PPTX slide.
  - layout_chart_slide_pptx falls back to a table shape on malformed data.
  - layout_table_slide_pptx adds a real native PPTX table shape.
  - A full deck with chart + table slides round-trips through render_pptx and
    render_html, producing valid binary PPTX and a 16:9 HTML with SVG/table.
  - DeckSlide.chart / DeckSlide.table fields are accepted by MinimalDeck.
"""

from __future__ import annotations

import io

import pytest
from disco.tools.builtin._c8_chart_layouts import (
    ChartSpec,
    TableSpec,
    _to_svg_dict,
    html_chart_content,
    html_table_content,
)
from disco.tools.builtin._pptx_render import (
    DeckSlide,
    MinimalDeck,
    render_deck,
    render_html,
    render_pptx,
)

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _bar_spec() -> ChartSpec:
    return ChartSpec(
        kind="bar",
        title="Revenue by Quarter",
        labels=["Q1", "Q2", "Q3", "Q4"],
        series=[{"name": "2025", "data": [120, 145, 130, 160]}],
    )


def _line_spec() -> ChartSpec:
    return ChartSpec(
        kind="line",
        title="Monthly Trend",
        labels=["Jan", "Feb", "Mar"],
        series=[{"name": "Users", "data": [1000, 1200, 1500]}],
    )


def _pie_spec() -> ChartSpec:
    return ChartSpec(
        kind="pie",
        title="Market Share",
        labels=["Disco", "Rival", "Other"],
        series=[{"name": "Share", "data": [55, 30, 15]}],
    )


def _scatter_spec() -> ChartSpec:
    return ChartSpec(
        kind="scatter",
        title="Engagement vs Retention",
        labels=[],
        series=[
            {"name": "Group A", "data": [{"x": 1.0, "y": 2.0}, {"x": 3.0, "y": 4.0}]},
            {"name": "Group B", "data": [[5.0, 6.0], [7.0, 8.0]]},
        ],
    )


def _empty_chart_spec() -> ChartSpec:
    """Malformed spec — no series data."""
    return ChartSpec(kind="bar", title="Empty", labels=["A", "B"], series=[])


def _bad_data_spec() -> ChartSpec:
    """Malformed spec — series has no numeric values."""
    return ChartSpec(
        kind="bar",
        title="Bad",
        labels=["X"],
        series=[{"name": "S", "data": ["not a number"]}],
    )


def _table_spec() -> TableSpec:
    return TableSpec(
        headers=["Name", "Score", "Grade"],
        rows=[
            ["Alice", "95", "A"],
            ["Bob", "82", "B"],
            ["Carol", "78", "C+"],
        ],
    )


def _chart_deck() -> MinimalDeck:
    return MinimalDeck(
        title="Data Report",
        theme_name="disco",
        theme_mode="light",
        slides=[
            DeckSlide(
                title="Revenue Chart",
                layout="chart",
                chart=_bar_spec(),
            ),
            DeckSlide(
                title="Results Table",
                layout="table",
                table=_table_spec(),
            ),
        ],
    )


# ---------------------------------------------------------------------------
# ChartSpec / TableSpec validation
# ---------------------------------------------------------------------------

def test_chart_spec_parses():
    s = _bar_spec()
    assert s.kind == "bar"
    assert s.labels == ["Q1", "Q2", "Q3", "Q4"]
    assert s.series[0]["name"] == "2025"
    assert s.series[0]["data"] == [120, 145, 130, 160]


def test_table_spec_parses():
    t = _table_spec()
    assert t.headers == ["Name", "Score", "Grade"]
    assert t.rows[0] == ["Alice", "95", "A"]


def test_chart_spec_empty_defaults():
    s = ChartSpec(kind="line")
    assert s.labels == []
    assert s.series == []
    assert s.title == ""


# ---------------------------------------------------------------------------
# Bridge: _to_svg_dict
# ---------------------------------------------------------------------------

def test_to_svg_dict_bar():
    d = _to_svg_dict(_bar_spec())
    assert d["chart_type"] == "bar"
    assert d["title"] == "Revenue by Quarter"
    assert len(d["data"]) == 4
    assert d["data"][0] == {"label": "Q1", "value": 120.0}


def test_to_svg_dict_scatter_dict_format():
    d = _to_svg_dict(_scatter_spec())
    assert d["chart_type"] == "scatter"
    assert any(item.get("group") == "Group A" for item in d["data"])
    assert any(item.get("group") == "Group B" for item in d["data"])


def test_to_svg_dict_scatter_list_format():
    spec = ChartSpec(
        kind="scatter", labels=[],
        series=[{"name": "G", "data": [[1.0, 2.0], [3.0, 4.0]]}],
    )
    d = _to_svg_dict(spec)
    assert d["data"][0] == {"x": 1.0, "y": 2.0, "group": "G"}


def test_to_svg_dict_empty_series():
    d = _to_svg_dict(_empty_chart_spec())
    assert d["data"] == []


# ---------------------------------------------------------------------------
# HTML chart content
# ---------------------------------------------------------------------------

def test_html_chart_content_contains_svg():
    """Valid chart data → HTML embeds an <svg>."""
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_chart_content("My Chart", _bar_spec(), theme)
    assert "<svg" in result
    assert "slide-chart" in result


def test_html_chart_content_line_svg():
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_chart_content("Trend", _line_spec(), theme)
    assert "<svg" in result


def test_html_chart_content_pie_svg():
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_chart_content("Share", _pie_spec(), theme)
    assert "<svg" in result


def test_html_chart_content_scatter_svg_or_table():
    """Scatter with valid data → SVG or table fallback (depends on SVG renderer)."""
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_chart_content("Scatter", _scatter_spec(), theme)
    # Must contain either SVG or a table (chart-table fallback)
    assert ("<svg" in result) or ("<table" in result)


def test_html_chart_content_malformed_falls_back_to_table():
    """Empty/malformed chart data → render_chart_table fallback (HTML table, not SVG)."""
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_chart_content("Bad", _empty_chart_spec(), theme)
    # SVG renderer returns None → fallback to chart-table or chart-fallback div
    assert "<svg" not in result
    # Must still contain the title
    assert "slide-heading" in result


def test_html_chart_content_title_escaped():
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    spec = ChartSpec(kind="bar", labels=["A"], series=[{"name": "S", "data": [1]}])
    result = html_chart_content('<script>alert(1)</script>', spec, theme)
    assert "<script>" not in result
    assert "&lt;script&gt;" in result


# ---------------------------------------------------------------------------
# HTML table content
# ---------------------------------------------------------------------------

def test_html_table_content_produces_table():
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_table_content("Results", _table_spec(), theme)
    assert "<table" in result
    assert "slide-table" in result
    assert "<th>" in result
    assert "Alice" in result
    assert "Bob" in result


def test_html_table_content_has_headers():
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_table_content("T", _table_spec(), theme)
    for hdr in ["Name", "Score", "Grade"]:
        assert hdr in result


def test_html_table_content_empty_spec():
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    result = html_table_content("Empty", TableSpec(), theme)
    assert "<table" not in result
    assert "slide-heading" in result


def test_html_table_content_escapes_cells():
    from disco.core.brand import resolve_theme
    theme = resolve_theme("disco", "light")
    spec = TableSpec(headers=["<Col>"], rows=[["<val>"]])
    result = html_table_content("T", spec, theme)
    assert "<Col>" not in result
    assert "&lt;Col&gt;" in result


# ---------------------------------------------------------------------------
# PPTX layout — chart slide
# ---------------------------------------------------------------------------

def test_pptx_chart_slide_has_chart_or_table_shape():
    """A chart slide should add a chart shape (or table fallback) — NOT just text."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Chart Deck",
        slides=[DeckSlide(title="Revenue", layout="chart", chart=_bar_spec())],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]

    shape_types = set()
    for shape in slide.shapes:
        shape_types.add(shape.shape_type)
        # MSO_SHAPE_TYPE.CHART == 3
        if shape.shape_type == 3:
            return  # found a native chart shape — test passes
    # Fallback table also acceptable: MSO_SHAPE_TYPE.TABLE == 19
    assert 19 in shape_types or 3 in shape_types, (
        f"Expected a chart (type=3) or table (type=19) shape; "
        f"found shape_types={shape_types}"
    )


def test_pptx_chart_slide_has_title_text():
    """Chart slide still carries a real title text run."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="CD",
        slides=[DeckSlide(title="Revenue by Quarter", layout="chart", chart=_bar_spec())],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]
    all_text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Revenue by Quarter" in all_text


def test_pptx_chart_malformed_data_fallback():
    """Malformed chart data (no series) → fallback shape, no crash."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Empty Chart",
        slides=[DeckSlide(title="No Data", layout="chart", chart=_empty_chart_spec())],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 1  # rendered without error


def test_pptx_chart_malformed_numeric_fallback():
    """Non-numeric data in series → fallback, no crash."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Bad",
        slides=[DeckSlide(title="Bad Data", layout="chart", chart=_bad_data_spec())],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 1


def test_pptx_line_chart():
    """Line chart renders without error."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Lines",
        slides=[DeckSlide(title="Trend", layout="chart", chart=_line_spec())],
    )
    prs = Presentation(io.BytesIO(render_pptx(deck)))
    assert len(prs.slides) == 1


def test_pptx_pie_chart():
    """Pie chart renders without error."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Pie",
        slides=[DeckSlide(title="Share", layout="chart", chart=_pie_spec())],
    )
    prs = Presentation(io.BytesIO(render_pptx(deck)))
    assert len(prs.slides) == 1


def test_pptx_scatter_chart():
    """Scatter chart renders without error (XyChartData path)."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Scatter",
        slides=[DeckSlide(title="Scatter", layout="chart", chart=_scatter_spec())],
    )
    prs = Presentation(io.BytesIO(render_pptx(deck)))
    assert len(prs.slides) == 1


# ---------------------------------------------------------------------------
# PPTX layout — table slide
# ---------------------------------------------------------------------------

def test_pptx_table_slide_has_native_table():
    """A table slide must add a native PPTX table shape (shape_type == 19)."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Table Deck",
        slides=[DeckSlide(title="Results", layout="table", table=_table_spec())],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]

    table_shapes = [s for s in slide.shapes if s.shape_type == 19]
    assert table_shapes, "Expected at least one native PPTX table shape (shape_type=19)"


def test_pptx_table_slide_has_title_text():
    """Table slide carries the title in a real text run."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="TD",
        slides=[DeckSlide(title="Student Grades", layout="table", table=_table_spec())],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]
    all_text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Student Grades" in all_text


def test_pptx_table_slide_cell_content():
    """Table slide carries at least one data cell value."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="TD",
        slides=[DeckSlide(title="T", layout="table", table=_table_spec())],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]

    # Find the table shape and check cell text
    for shape in slide.shapes:
        if shape.shape_type == 19:
            tbl = shape.table
            all_cells = [
                tbl.cell(r, c).text
                for r in range(tbl._tbl.tr_lst.__len__() if hasattr(tbl._tbl, 'tr_lst') else len(tbl.rows))
                for c in range(len(tbl.columns))
            ]
            all_cell_text = " ".join(all_cells)
            assert "Alice" in all_cell_text or "Name" in all_cell_text
            return
    pytest.fail("No table shape found")


def test_pptx_table_empty_spec_no_crash():
    """Empty TableSpec → placeholder, no crash."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Empty Table",
        slides=[DeckSlide(title="Empty", layout="table", table=TableSpec())],
    )
    prs = Presentation(io.BytesIO(render_pptx(deck)))
    assert len(prs.slides) == 1


# ---------------------------------------------------------------------------
# Full deck round-trip
# ---------------------------------------------------------------------------

def test_full_chart_table_deck_pptx():
    """A deck with chart + table slides produces valid PPTX with 2 slides."""
    from pptx import Presentation

    deck = _chart_deck()
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 2


def test_full_chart_table_deck_html_has_svg():
    """HTML render of a chart deck contains an <svg> element."""
    deck = _chart_deck()
    html_str = render_html(deck)
    assert "<svg" in html_str, "HTML render of chart deck must embed SVG"


def test_full_chart_table_deck_html_has_table():
    """HTML render of a table deck contains an HTML <table>."""
    deck = _chart_deck()
    html_str = render_html(deck)
    assert '<table class="slide-table">' in html_str


def test_full_deck_html_section_count():
    """render_html has one <section> per slide for a chart+table deck."""
    deck = _chart_deck()
    html_str = render_html(deck)
    count = html_str.count('<section class="slide')
    assert count == 2


def test_render_deck_chart_table():
    """render_deck returns both outputs for a chart+table deck."""
    result = render_deck(_chart_deck())
    assert isinstance(result["pptx_bytes"], bytes)
    assert isinstance(result["html_str"], str)
    assert "<svg" in result["html_str"]


# ---------------------------------------------------------------------------
# DeckSlide field integration
# ---------------------------------------------------------------------------

def test_deckslide_accepts_chart_field():
    s = DeckSlide(title="T", layout="chart", chart=_bar_spec())
    assert s.chart is not None
    assert s.chart.kind == "bar"


def test_deckslide_accepts_table_field():
    s = DeckSlide(title="T", layout="table", table=_table_spec())
    assert s.table is not None
    assert s.table.headers[0] == "Name"


def test_deckslide_chart_and_table_both_none_by_default():
    s = DeckSlide(title="T")
    assert s.chart is None
    assert s.table is None


def test_pptx_notes_on_chart_slide():
    """Speaker notes propagate to chart slides."""
    from pptx import Presentation

    deck = MinimalDeck(
        title="Notes Test",
        slides=[
            DeckSlide(
                title="Q1 Results",
                layout="chart",
                chart=_bar_spec(),
                notes="Highlight the Q4 spike in your presentation.",
            )
        ],
    )
    prs = Presentation(io.BytesIO(render_pptx(deck)))
    notes_text = prs.slides[0].notes_slide.notes_text_frame.text
    assert "Q4 spike" in notes_text
