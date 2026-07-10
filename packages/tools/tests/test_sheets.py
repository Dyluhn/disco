"""Unit tests for the sheet_generate tool.

Proves:
  - round-trip (write → reopen with openpyxl, assert formulas preserved as
    STRINGS not evaluated values);
  - whitelist enforcement (allowed formulas pass; IMPORTXML/INDIRECT/WEBSERVICE/
    HYPERLINK, unknown functions, and injection triggers become inert text);
  - the write is JAILED: it goes through the sandbox (path-escape attempts are
    rejected) and the artifact lands in the workspace — NOT the host cwd.

The tool declares runs_in="sandbox", so every test drives it through a REAL
process-backend sandbox instance jailed to a temp workspace. That backend's
_resolve() rejects ../ and absolute escapes — the same clamp production uses.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest
from disco.core.llm import ModelExecutionPolicy
from disco.tools.anatomy import ToolContext
from disco.tools.builtin import build_default_registry
from disco.tools.builtin.sheets import SheetGenerateArgs, SheetsTool
from disco.tools.executor import DefaultToolExecutor
from disco.tools.registry import agent_scope
from disco.tools.sandbox.base import SandboxSpec
from disco.tools.sandbox.process import ProcessSandboxInstance
from disco.tools.secrets import CapabilityBroker
from tool_fakes import FakeSandboxInstance, call


@pytest.fixture
def tmp_workspace():
    """A real temp workspace dir that the jailed sandbox writes into."""
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


def _jailed_sandbox(workspace: Path) -> ProcessSandboxInstance:
    """The real process backend, jailed to `workspace` — its write_file rejects
    path escapes (../ and absolute), exactly as production does."""
    return ProcessSandboxInstance(
        id="test-sbx",
        owner_id="test",
        conversation_id="test-cid",
        spec=SandboxSpec(),
        workspace=workspace,
    )


def _ctx(sandbox) -> ToolContext:
    return ToolContext(
        sandbox=sandbox,
        workspace_path=".",  # jailed: relative to the sandbox workspace
        timeout_s=30,
        capabilities=CapabilityBroker().grant(frozenset()),
        owner_id="test",
        conversation_id="test-cid",
    )


def _assert_escaped_literal(cell, raw: str) -> None:
    assert cell.data_type == "s"
    assert cell.value == "'" + raw


async def test_round_trip_formulas_preserved_as_strings(tmp_workspace):
    """Write a workbook with formulas, reopen with openpyxl, assert:
    - Formulas are preserved as STRINGS (e.g. '=SUM(A2:A4)'), not evaluated values.
    - Values are stored correctly.
    - Column headers are present.
    """
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Test Workbook",
        filename="formula_test.xlsx",
        sheets=[
            dict(
                name="Sales",
                columns=["Item", "Price", "Qty", "Total"],
                rows=[
                    ["Widget", 10, 5, "=PRODUCT(B2,C2)"],
                    ["Gadget", 15, 3, "=PRODUCT(B3,C3)"],
                    ["Doohickey", 7, 8, "=PRODUCT(B4,C4)"],
                    ["SUBTOTAL", "", "", "=SUM(D2:D4)"],
                ],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.error}"
    assert outcome.artifacts == ["formula_test.xlsx"]

    # The file landed IN the jailed workspace (not the host cwd).
    wb = openpyxl.load_workbook(tmp_workspace / "formula_test.xlsx")
    ws = wb["Sales"]

    # Column headers.
    assert ws.cell(row=1, column=1).value == "Item"
    assert ws.cell(row=1, column=2).value == "Price"

    # Row 2: Widget, 10, 5, =PRODUCT(B2,C2)
    assert ws.cell(row=2, column=1).value == "Widget"
    assert ws.cell(row=2, column=2).value == 10
    assert ws.cell(row=2, column=3).value == 5
    # KEY ASSERTION: allowed formula preserved, NOT the evaluated value 50.
    assert ws.cell(row=2, column=4).value == "=PRODUCT(B2,C2)", (
        f"Expected formula string '=PRODUCT(B2,C2)', got {ws.cell(row=2, column=4).value!r}"
    )

    # Row 5: SUM total — formula preserved as '=SUM(D2:D4)', NOT 50+45+56=151.
    assert ws.cell(row=5, column=4).value == "=SUM(D2:D4)", (
        f"Expected formula string '=SUM(D2:D4)', got {ws.cell(row=5, column=4).value!r}"
    )

    wb.close()


def test_sheet_rows_schema_advertises_scalar_cell_types():
    schema = SheetsTool.definition.to_spec().parameters_schema
    cell_schema = schema["properties"]["sheets"]["items"]["properties"]["rows"]["items"]["items"]
    advertised_types = {branch["type"] for branch in cell_schema["anyOf"]}

    assert advertised_types == {"string", "number", "integer", "boolean", "null"}


async def test_boolean_and_null_cells_round_trip(tmp_workspace):
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Flags",
        filename="flags.xlsx",
        sheets=[
            dict(
                name="Flags",
                columns=["Name", "Enabled", "Notes"],
                rows=[
                    ["Alpha", True, None],
                    ["Beta", False, "paused"],
                ],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.error}"

    wb = openpyxl.load_workbook(tmp_workspace / "flags.xlsx")
    ws = wb["Flags"]
    assert ws.cell(row=2, column=2).value is True
    assert ws.cell(row=2, column=3).value is None
    assert ws.cell(row=3, column=2).value is False
    wb.close()


async def test_formula_whitelist_escapes_importxml(tmp_workspace):
    """IMPORTXML is never written as a live formula."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Bad Workbook",
        filename="bad.xlsx",
        sheets=[
            dict(
                name="Sheet1",
                columns=["A"],
                rows=[['=IMPORTXML("http://evil.com", "//data")']],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    wb = openpyxl.load_workbook(tmp_workspace / "bad.xlsx")
    _assert_escaped_literal(
        wb["Sheet1"].cell(row=2, column=1),
        '=IMPORTXML("http://evil.com", "//data")',
    )
    wb.close()


async def test_formula_whitelist_escapes_indirect(tmp_workspace):
    """INDIRECT is never written as a live formula."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="INDIRECT Attempt",
        filename="indirect.xlsx",
        sheets=[
            dict(
                name="Sheet1",
                columns=["A"],
                rows=[['=INDIRECT("B" & ROW())']],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    wb = openpyxl.load_workbook(tmp_workspace / "indirect.xlsx")
    _assert_escaped_literal(wb["Sheet1"].cell(row=2, column=1), '=INDIRECT("B" & ROW())')
    wb.close()


async def test_whitespace_function_call_still_validated(tmp_workspace):
    """A function call with whitespace before '(' — `INDIRECT ("A1")` — must STILL
    be neutralized. Some spreadsheet engines tolerate that spacing; if the matcher
    anchored '(' immediately after the name, the call would slip past the
    deny-list entirely. This is the regression guard for that bypass."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Spaced INDIRECT",
        filename="spaced.xlsx",
        sheets=[
            dict(
                name="Sheet1",
                columns=["A"],
                rows=[['=INDIRECT ("A1")']],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    wb = openpyxl.load_workbook(tmp_workspace / "spaced.xlsx")
    _assert_escaped_literal(wb["Sheet1"].cell(row=2, column=1), '=INDIRECT ("A1")')
    wb.close()


async def test_formula_whitelist_escapes_webservice(tmp_workspace):
    """WEBSERVICE is never written as a live formula."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="WEBSERVICE Attempt",
        filename="webservice.xlsx",
        sheets=[
            dict(
                name="Sheet1",
                columns=["A"],
                rows=[['=WEBSERVICE("http://example.com/api")']],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    wb = openpyxl.load_workbook(tmp_workspace / "webservice.xlsx")
    _assert_escaped_literal(
        wb["Sheet1"].cell(row=2, column=1),
        '=WEBSERVICE("http://example.com/api")',
    )
    wb.close()


async def test_formula_whitelist_escapes_hyperlink(tmp_workspace):
    """HYPERLINK is never written as a live formula."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="HYPERLINK Attempt",
        filename="hyperlink.xlsx",
        sheets=[
            dict(
                name="Sheet1",
                columns=["A"],
                rows=[['=HYPERLINK("http://phishing.example.com", "Click here")']],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    wb = openpyxl.load_workbook(tmp_workspace / "hyperlink.xlsx")
    _assert_escaped_literal(
        wb["Sheet1"].cell(row=2, column=1),
        '=HYPERLINK("http://phishing.example.com", "Click here")',
    )
    wb.close()


async def test_allowed_formulas_pass(tmp_workspace):
    """A sheet with diverse allowed formulas (SUM, AVERAGE, IF, VLOOKUP, INDEX,
    MATCH, DATE) passes validation and writes successfully."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Allowed Formulas",
        filename="allowed.xlsx",
        sheets=[
            dict(
                name="Data",
                columns=["Category", "Value", "Comment"],
                rows=[
                    ["A", 100, "=SUM(B2:B5)"],
                    ["B", 200, "=AVERAGE(B2:B5)"],
                    ["C", 150, '=IF(B3>100, "HIGH", "LOW")'],
                    ["D", 50, '=VLOOKUP("A", A2:B5, 2, FALSE)'],
                    ["E", "", '=INDEX(B2:B5, MATCH("C", A2:A5, 0))'],
                    ["Date", "", "=DATE(2026, 6, 11)"],
                ],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    assert outcome.artifacts == ["allowed.xlsx"]

    wb = openpyxl.load_workbook(tmp_workspace / "allowed.xlsx")
    ws = wb["Data"]

    assert ws.cell(row=2, column=3).value == "=SUM(B2:B5)"
    assert ws.cell(row=3, column=3).value == "=AVERAGE(B2:B5)"
    assert ws.cell(row=4, column=3).value == '=IF(B3>100, "HIGH", "LOW")'
    assert ws.cell(row=7, column=3).value == "=DATE(2026, 6, 11)"

    wb.close()


async def test_unknown_formula_is_escaped(tmp_workspace):
    """A function outside the whitelist is never written live."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Unknown Fn",
        filename="unknown.xlsx",
        sheets=[
            dict(
                name="Sheet1",
                columns=["A"],
                rows=[["=FOOBARBAZ(A1:A10)"]],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    wb = openpyxl.load_workbook(tmp_workspace / "unknown.xlsx")
    _assert_escaped_literal(wb["Sheet1"].cell(row=2, column=1), "=FOOBARBAZ(A1:A10)")
    wb.close()


async def test_formula_injection_triggers_are_inert_in_headers_and_rows(tmp_workspace):
    """Exercise DDE, external references, leading sigils/whitespace, a bare
    reference expression, and an allowed first function followed by a blocked
    token. Every payload must survive only as an explicit string cell."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))
    malicious = [
        "=cmd|'/c calc'!A1",
        "=SUM(A1:A2)|cmd",
        "=SUM([1]Sheet1!A1)",
        "=1+1",
        "=A1+B1",
        '+HYPERLINK("http://evil.example", "x")',
        "-2+3",
        "@SUM(1)",
        "\t=SUM(1)",
        "[external.xlsx]Sheet1",
        "!Sheet1",
        "|cmd",
    ]
    args = SheetGenerateArgs(
        title="Injection Literals",
        filename="injection.xlsx",
        sheets=[dict(name="Payloads", columns=malicious, rows=[malicious])],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"
    wb = openpyxl.load_workbook(tmp_workspace / "injection.xlsx")
    ws = wb["Payloads"]
    for column, raw in enumerate(malicious, start=1):
        _assert_escaped_literal(ws.cell(row=1, column=column), raw)
        _assert_escaped_literal(ws.cell(row=2, column=column), raw)
    wb.close()


async def test_path_traversal_filename_rejected(tmp_workspace):
    """A model-controlled filename that escapes the workspace (../) must be
    rejected by the sandbox jail and surfaced as a clean failure — and NOTHING
    is written outside the workspace. This is the regression guard for the
    host-cwd write bug (the tool used to wb.save() directly to ctx.workspace_path
    '.' on the host, bypassing the jail)."""
    reg = build_default_registry()
    ex = DefaultToolExecutor(
        reg,
        agent_scope(model_policy=ModelExecutionPolicy.standard()),
        sandbox=_jailed_sandbox(tmp_workspace),
    )

    escape_target = tmp_workspace.parent / "escape.xlsx"
    assert not escape_target.exists()

    res = await ex.execute(
        call(
            "sheet_generate",
            title="Escape",
            filename="../escape.xlsx",
            sheets=[dict(name="S", columns=["A"], rows=[["x"]])],
        )
    )

    assert res.success is False, "path escape must fail, not write outside the jail"
    assert not escape_target.exists(), "no file may be written outside the workspace"


async def test_multiple_sheets(tmp_workspace):
    """Workbook with multiple sheets is written correctly."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Multi-Sheet",
        filename="multi.xlsx",
        sheets=[
            dict(
                name="Revenue",
                columns=["Month", "Amount"],
                rows=[["Jan", 1000], ["Feb", 1200], ["Total", "=SUM(B2:B3)"]],
            ),
            dict(
                name="Costs",
                columns=["Month", "Amount"],
                rows=[["Jan", 600], ["Feb", 700], ["Total", "=SUM(B2:B3)"]],
            ),
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success, f"Tool failed: {outcome.content}"

    wb = openpyxl.load_workbook(tmp_workspace / "multi.xlsx")
    assert set(wb.sheetnames) == {"Revenue", "Costs"}
    assert wb["Revenue"].cell(row=4, column=2).value == "=SUM(B2:B3)"
    assert wb["Costs"].cell(row=4, column=2).value == "=SUM(B2:B3)"
    wb.close()


async def test_tool_registered_and_scoped():
    """sheet_generate is in the default registry AND in agent scope."""
    reg = build_default_registry()
    agt_scope = agent_scope(model_policy=ModelExecutionPolicy.standard())
    assert "sheet_generate" in reg.names()
    assert "sheet_generate" in agt_scope.allowed_tools

    tool = reg.get("sheet_generate", scope=agt_scope)
    assert tool is not None
    assert tool.definition.name == "sheet_generate"


async def test_executor_rejects_sheet_generate_in_research_scope():
    """sheet_generate is agent-scope only, not research."""
    from disco.tools.registry import research_scope

    reg = build_default_registry()
    ex = DefaultToolExecutor(reg, research_scope(), sandbox=FakeSandboxInstance())
    res = await ex.execute(call("sheet_generate", title="Test", filename="x.xlsx", sheets=[]))
    assert res.success is False
    assert res.structured["kind"] == "unknown_tool"


async def test_simple_value_sheet(tmp_workspace):
    """A sheet with plain values only (no formulas) writes correctly."""
    tool = SheetsTool()
    ctx = _ctx(_jailed_sandbox(tmp_workspace))

    args = SheetGenerateArgs(
        title="Plain Values",
        filename="plain.xlsx",
        sheets=[
            dict(
                name="Data",
                columns=["Name", "Age", "City"],
                rows=[
                    ["Alice", 30, "NYC"],
                    ["Bob", 25, "LA"],
                    ["Carol", 35, "Chicago"],
                ],
            )
        ],
    )

    outcome = await tool.run(args, ctx)
    assert outcome.success

    wb = openpyxl.load_workbook(tmp_workspace / "plain.xlsx")
    ws = wb["Data"]
    assert ws.cell(row=2, column=1).value == "Alice"
    assert ws.cell(row=2, column=2).value == 30
    assert ws.cell(row=3, column=3).value == "LA"
    wb.close()
