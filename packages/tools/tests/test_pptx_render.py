"""Tests for C3 — native editable PPTX renderer (_pptx_render.py).

Proves:
  - render_pptx: the .pptx round-trips via python-pptx and contains REAL text
    runs (NOT a single embedded image), one shape per layout element.
  - render_html: the output is valid 16:9 HTML with one <section> per slide,
    inlined brand fonts, and keyboard-nav script.
  - convert_to_pdf: soffice-absent returns a clear failure (no false affordance).
  - render_deck: returns both pptx_bytes and html_str.
  - Binary round-trip: pptx_bytes are valid ZIP/OOXML (magic bytes 50 4B).
  - Each layout variant (title / bullets / section / image_right) renders
    without error and deposits text in the slide.
"""

from __future__ import annotations

import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import pytest
from disco.core.contract.export_render import check_export_render, pptx_visible_text
from disco.tools.builtin._pptx_render import (
    DeckSlide,
    MinimalDeck,
    convert_to_pdf,
    render_deck,
    render_html,
    render_pptx,
)

# ---------------------------------------------------------------------------
# Sample decks
# ---------------------------------------------------------------------------


def _sample_deck() -> MinimalDeck:
    """A minimal multi-slide deck exercising all four layouts."""
    return MinimalDeck(
        title="Test Deck",
        theme_name="disco",
        theme_mode="light",
        slides=[
            DeckSlide(
                title="Welcome to Disco",
                bullets=["The native PPTX renderer"],
                layout="title",
            ),
            DeckSlide(
                title="Key Points",
                bullets=["Real text boxes", "Brand fonts", "16:9 canvas"],
                layout="bullets",
            ),
            DeckSlide(
                title="Deep Dive",
                bullets=["A section subtitle"],
                layout="section",
            ),
            DeckSlide(
                title="Image + Bullets",
                bullets=["With placeholder", "For C7 wiring"],
                layout="image_right",
                image_url=None,
            ),
        ],
    )


def _dark_deck() -> MinimalDeck:
    """A single-slide disco-dark deck."""
    return MinimalDeck(
        title="Dark Mode",
        theme_name="disco",
        theme_mode="dark",
        slides=[
            DeckSlide(title="Dark Slide", bullets=["Dark theme"], layout="bullets"),
        ],
    )


def _neutral_deck() -> MinimalDeck:
    """A neutral-theme deck."""
    return MinimalDeck(
        title="Neutral",
        theme_name="neutral",
        theme_mode="light",
        slides=[
            DeckSlide(title="Neutral Slide", bullets=["System fonts"], layout="bullets"),
        ],
    )


# ---------------------------------------------------------------------------
# PPTX tests
# ---------------------------------------------------------------------------


def test_pptx_binary_magic_bytes():
    """The pptx_bytes start with the ZIP magic bytes PK (50 4B 03 04)."""
    deck = _sample_deck()
    data = render_pptx(deck)
    assert isinstance(data, bytes)
    assert data[:4] == b"PK\x03\x04", "Expected OOXML ZIP magic bytes"


def test_pptx_round_trips_via_python_pptx():
    """python-pptx can re-open the rendered bytes without error."""
    import io

    from pptx import Presentation

    deck = _sample_deck()
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    # Should have same slide count
    assert len(prs.slides) == len(deck.slides)


def test_pptx_contains_real_text_runs():
    """The .pptx contains real text runs — NOT a single embedded image per slide.

    Specifically: at least one TextFrame with the slide title text must be
    present for each slide.  This proves text boxes, not image shapes.
    """
    import io

    from pptx import Presentation
    from pptx.shapes.picture import Picture

    deck = _sample_deck()
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))

    for i, (pptx_slide, deck_slide) in enumerate(zip(prs.slides, deck.slides, strict=True)):
        # Collect all text in the slide
        all_text = ""
        pic_count = 0
        for shape in pptx_slide.shapes:
            if shape.has_text_frame:
                all_text += shape.text_frame.text
            if isinstance(shape, Picture):
                pic_count += 1

        assert deck_slide.title in all_text, (
            f"Slide {i} ({deck_slide.layout!r}): title {deck_slide.title!r} "
            f"not found in extracted text {all_text!r}"
        )
        # Image-per-slide would be ONE Picture shape with no title text.
        # We should have zero pictures on non-image_right layouts.
        if deck_slide.layout != "image_right":
            assert pic_count == 0, (
                f"Slide {i} ({deck_slide.layout!r}): unexpected Picture shape — "
                "renderer must use text boxes, not images"
            )


def test_pptx_16_9_canvas():
    """Slide dimensions are 12 192 000 × 6 858 000 EMU (16:9)."""
    import io

    from pptx import Presentation

    deck = _sample_deck()
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    assert prs.slide_width == 12_192_000
    assert prs.slide_height == 6_858_000


def test_pptx_bullets_slide_has_bullet_text():
    """A bullets slide carries each bullet in at least one text run."""
    import io

    from pptx import Presentation

    deck = MinimalDeck(
        title="Bullets Test",
        slides=[
            DeckSlide(
                title="Three Bullets",
                bullets=["Alpha", "Beta", "Gamma"],
                layout="bullets",
            )
        ],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]
    all_text = " ".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
    for bullet in ("Alpha", "Beta", "Gamma"):
        assert bullet in all_text, f"Bullet {bullet!r} missing from slide text"


def test_pptx_title_layout():
    """Title layout renders the deck title and optional subtitle."""
    import io

    from pptx import Presentation

    deck = MinimalDeck(
        title="Title Test",
        slides=[
            DeckSlide(
                title="Grand Title",
                bullets=["A subtitle here"],
                layout="title",
            )
        ],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]
    all_text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Grand Title" in all_text
    assert "A subtitle here" in all_text


def test_pptx_section_layout():
    """Section layout renders the section title."""
    import io

    from pptx import Presentation

    deck = MinimalDeck(
        title="Section Test",
        slides=[
            DeckSlide(title="Chapter One", layout="section"),
        ],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]
    all_text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Chapter One" in all_text


def test_pptx_dark_theme():
    """Dark theme renders without error and produces valid bytes."""
    import io

    from pptx import Presentation

    data = render_pptx(_dark_deck())
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 1


def test_pptx_neutral_theme():
    """Neutral theme renders without error."""
    import io

    from pptx import Presentation

    data = render_pptx(_neutral_deck())
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 1


def test_pptx_with_notes():
    """Speaker notes land in the notes slide text frame."""
    import io

    from pptx import Presentation

    deck = MinimalDeck(
        title="Notes Test",
        slides=[
            DeckSlide(
                title="Noted Slide",
                bullets=["Bullet"],
                layout="bullets",
                notes="These are the speaker notes.",
            )
        ],
    )
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    notes_text = prs.slides[0].notes_slide.notes_text_frame.text
    assert "speaker notes" in notes_text


def test_pptx_empty_deck():
    """An empty deck (no slides) produces valid 0-slide PPTX."""
    import io

    from pptx import Presentation

    deck = MinimalDeck(title="Empty", slides=[])
    data = render_pptx(deck)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 0


# ---------------------------------------------------------------------------
# HTML tests
# ---------------------------------------------------------------------------


def test_html_is_string():
    """render_html returns a string."""
    html_str = render_html(_sample_deck())
    assert isinstance(html_str, str)


def test_html_one_section_per_slide():
    """The HTML contains one <section class="slide"> per slide."""
    deck = _sample_deck()
    html_str = render_html(deck)
    count = html_str.count('<section class="slide')
    assert count == len(deck.slides), f"Expected {len(deck.slides)} slide sections, found {count}"


def test_html_16_9_ratio():
    """The HTML shell uses 56.25vw height (16:9 via viewport units)."""
    html_str = render_html(_sample_deck())
    assert "56.25vw" in html_str, "Missing 16:9 viewport-unit height in CSS"


def test_html_contains_title_text():
    """Each slide title appears in the HTML."""
    deck = _sample_deck()
    html_str = render_html(deck)
    for slide in deck.slides:
        assert slide.title in html_str, f"Slide title {slide.title!r} not found in HTML"


def test_html_has_keyboard_nav_script():
    """The HTML includes a keyboard navigation script (← / →)."""
    html_str = render_html(_sample_deck())
    assert "ArrowRight" in html_str
    assert "ArrowLeft" in html_str


def test_html_has_font_face_declarations():
    """The HTML inlines @font-face declarations for brand OFL fonts."""
    html_str = render_html(_sample_deck())
    # All three brand typefaces must be declared
    assert "@font-face" in html_str
    assert "Fraunces" in html_str
    assert "Schibsted Grotesk" in html_str
    assert "Newsreader" in html_str


def test_html_has_brand_css_vars():
    """Brand CSS custom properties are present in the HTML."""
    html_str = render_html(_sample_deck())
    assert "--bg:" in html_str
    assert "--accent:" in html_str
    assert "--text:" in html_str


def test_html_first_slide_is_active():
    """The first slide has class 'active'; others do not on initial load."""
    html_str = render_html(_sample_deck())
    # The first active marker
    first_active = html_str.find('class="slide active"')
    assert first_active != -1, "First slide must have 'active' class"
    # Only one initial active slide
    assert html_str.count('class="slide active"') == 1


def test_html_escapes_special_chars():
    """Slide titles with HTML special chars are properly escaped."""
    deck = MinimalDeck(
        title="<Test & Check>",
        slides=[
            DeckSlide(
                title='<Script> "injection"',
                bullets=["<b>bold attempt</b>"],
                layout="bullets",
            )
        ],
    )
    html_str = render_html(deck)
    # Raw angle brackets from title/bullet must not appear unescaped
    assert "<Script>" not in html_str
    assert "<b>bold" not in html_str
    # Escaped versions must be present
    assert "&lt;Script&gt;" in html_str
    assert "&lt;b&gt;" in html_str


def test_html_dark_theme_bg():
    """Dark theme background color appears in the HTML CSS vars."""
    html_str = render_html(_dark_deck())
    # Disco dark bg is #0e0f12
    assert "#0e0f12" in html_str


# ---------------------------------------------------------------------------
# PDF / convert_to_pdf tests
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_convert_to_pdf_absent_soffice_returns_clear_failure():
    """When soffice is absent, convert_to_pdf returns (False, descriptive_error)."""
    ctx = MagicMock()
    mock_result = MagicMock()
    mock_result.exit_code = 1
    ctx.sandbox.exec_shell = AsyncMock(return_value=mock_result)

    ok, err = await convert_to_pdf(ctx, "deck.pptx")

    assert ok is False
    assert err  # non-empty error message
    # Must mention soffice/LibreOffice — no silent failure
    assert "soffice" in err.lower() or "libreoffice" in err.lower()


@pytest.mark.asyncio
async def test_convert_to_pdf_probe_called_first():
    """convert_to_pdf probes for soffice before attempting conversion."""
    ctx = MagicMock()
    # First call (probe) → soffice absent
    probe_result = MagicMock(exit_code=1)
    ctx.sandbox.exec_shell = AsyncMock(return_value=probe_result)

    ok, err = await convert_to_pdf(ctx, "deck.pptx")

    assert ok is False
    # Only the probe call should have fired (no conversion attempt)
    assert ctx.sandbox.exec_shell.call_count == 1
    probe_call_args = ctx.sandbox.exec_shell.call_args_list[0]
    assert "command -v soffice" in str(probe_call_args)


@pytest.mark.asyncio
async def test_convert_to_pdf_success_path():
    """When soffice is present and conversion succeeds, returns (True, '')."""
    ctx = MagicMock()
    # Probe → found; conversion → success
    probe_ok = MagicMock(exit_code=0)
    convert_ok = MagicMock(exit_code=0, timed_out=False, stderr="")
    ctx.sandbox.exec_shell = AsyncMock(side_effect=[probe_ok, convert_ok])

    ok, err = await convert_to_pdf(ctx, "deck.pptx")

    assert ok is True
    assert err == ""
    # Second call should include soffice + pptx name
    convert_call = ctx.sandbox.exec_shell.call_args_list[1]
    assert "soffice" in str(convert_call)
    assert "deck.pptx" in str(convert_call)


@pytest.mark.asyncio
async def test_convert_to_pdf_timeout_returns_clear_failure():
    """A timed-out soffice returns (False, descriptive error)."""
    ctx = MagicMock()
    probe_ok = MagicMock(exit_code=0)
    timeout_result = MagicMock(exit_code=1, timed_out=True, stderr="")
    ctx.sandbox.exec_shell = AsyncMock(side_effect=[probe_ok, timeout_result])

    ok, err = await convert_to_pdf(ctx, "deck.pptx")

    assert ok is False
    assert "timed out" in err.lower()


@pytest.mark.asyncio
async def test_convert_to_pdf_nonzero_exit_returns_failure():
    """A non-zero soffice exit code returns (False, error text)."""
    ctx = MagicMock()
    probe_ok = MagicMock(exit_code=0)
    fail_result = MagicMock(exit_code=2, timed_out=False, stderr="Error: import failed")
    ctx.sandbox.exec_shell = AsyncMock(side_effect=[probe_ok, fail_result])

    ok, err = await convert_to_pdf(ctx, "deck.pptx")

    assert ok is False
    assert "import failed" in err or err  # either stderr or fallback


# ---------------------------------------------------------------------------
# render_deck tests
# ---------------------------------------------------------------------------


def test_render_deck_returns_both_outputs():
    """render_deck returns pptx_bytes (bytes) and html_str (str)."""
    result = render_deck(_sample_deck())
    assert "pptx_bytes" in result
    assert "html_str" in result
    assert isinstance(result["pptx_bytes"], bytes)
    assert isinstance(result["html_str"], str)


def test_render_deck_pptx_is_valid():
    """render_deck pptx_bytes are a valid OOXML file."""
    import io

    from pptx import Presentation

    result = render_deck(_sample_deck())
    prs = Presentation(io.BytesIO(result["pptx_bytes"]))
    assert len(prs.slides) == len(_sample_deck().slides)


def test_render_deck_html_is_16_9():
    """render_deck html_str is a 16:9 HTML deck."""
    result = render_deck(_sample_deck())
    assert "56.25vw" in result["html_str"]


# ---------------------------------------------------------------------------
# Default-template brand chrome (mirrors the PDF report: running wordmark on
# every slide + the definition-mark colophon on the cover). Gated on
# theme.branded.
# ---------------------------------------------------------------------------

from disco.tools.builtin._deck_schema import (  # noqa: E402
    AuthoredDeck,
    AuthoredSlide,
    lower_deck,
)


def _branded_deck(theme: str = "disco-light") -> AuthoredDeck:
    return AuthoredDeck(
        title="Brand Test",
        theme=theme,
        slides=[
            AuthoredSlide(type="title", title="Cover", body=["A subtitle"]),
            AuthoredSlide(type="bullets", title="Content", body=["one", "two"]),
            AuthoredSlide(type="section_header", title="Divider", body=["lead"]),
        ],
    )


def _pptx_slide_texts(pptx_bytes: bytes) -> list[list[str]]:
    import io

    from pptx import Presentation

    prs = Presentation(io.BytesIO(pptx_bytes))
    out = []
    for s in prs.slides:
        out.append([sh.text_frame.text for sh in s.shapes if sh.has_text_frame])
    return out


def test_pptx_wordmark_on_every_branded_slide():
    """The 'Disco' running wordmark must appear on EVERY slide of a branded deck."""
    texts = _pptx_slide_texts(render_pptx(lower_deck(_branded_deck())))
    assert len(texts) == 3
    for slide_texts in texts:
        assert any("Disco" in t for t in slide_texts), slide_texts


def test_pptx_colophon_only_on_cover():
    """The 'disco — Latin·verb' colophon rides ONLY the cover/title slide."""
    texts = _pptx_slide_texts(render_pptx(lower_deck(_branded_deck())))
    cover = " ".join(texts[0])
    assert "acquainted with" in cover and "discere" in cover
    for slide_texts in texts[1:]:
        joined = " ".join(slide_texts)
        assert "acquainted with" not in joined


def test_pptx_no_brand_marks_when_unbranded():
    """The neutral theme is unbranded → no wordmark, no colophon (clean canvas)."""
    texts = _pptx_slide_texts(render_pptx(lower_deck(_branded_deck(theme="neutral"))))
    flat = " ".join(t for slide in texts for t in slide)
    assert "Disco" not in flat
    assert "acquainted with" not in flat


def test_html_brand_marks_branded_vs_neutral():
    """HTML mirrors PPTX: brand-wordmark on every slide, brand-colophon on the
    cover only; neither under the neutral (unbranded) theme."""
    html_branded = render_html(lower_deck(_branded_deck()))
    # one wordmark div per slide (3 slides); match the element, not the CSS rule
    assert html_branded.count('class="brand-wordmark"') == 3
    # exactly one colophon div (the cover)
    assert html_branded.count('class="brand-colophon"') == 1
    assert "I become acquainted with" in html_branded

    html_neutral = render_html(lower_deck(_branded_deck(theme="neutral")))
    assert 'class="brand-wordmark"' not in html_neutral
    assert 'class="brand-colophon"' not in html_neutral


def test_html_brand_chrome_is_pointer_events_none():
    """Brand chrome must not intercept editor clicks — the SelectionOverlay resolves
    from event.target upward, and these nodes carry no data-element-id. pointer-events:
    none lets clicks fall through to the underlying slide element."""
    css = render_html(lower_deck(_branded_deck()))
    # Both brand layers opt out of pointer events.
    assert ".brand-wordmark{" in css.replace("\n", "") or "brand-wordmark" in css
    # The rule block carries pointer-events:none for both classes.
    assert css.count("pointer-events:none") >= 2


def _blank_branded_checker_deck() -> AuthoredDeck:
    return AuthoredDeck(
        title=" ",
        theme="disco-light",
        slides=[
            AuthoredSlide(type="title", title="   ", body=[]),
            AuthoredSlide(type="bullets", title=" ", body=["", "  "]),
            AuthoredSlide(type="bullets", title="", body=[]),
        ],
    )


def _content_branded_checker_deck() -> AuthoredDeck:
    return AuthoredDeck(
        title="Market Expansion",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="title",
                title="Market Expansion",
                body=["North America launch plan"],
            ),
            AuthoredSlide(
                type="bullets",
                title="Audience Signals",
                body=["Small teams need faster onboarding", "Buyers compare total cost"],
            ),
            AuthoredSlide(
                type="bullets",
                title="Launch Plan",
                body=["Pilot with five accounts", "Measure activation weekly"],
            ),
        ],
    )


def test_pptx_export_render_refuses_blank_branded_real_render() -> None:
    data = render_pptx(lower_deck(_blank_branded_checker_deck()))
    f = check_export_render("pptx", data, declared_units=3)
    assert f.non_blank is False and f.ok is False


@dataclass(frozen=True)
class _LocalExecResult:
    exit_code: int
    timed_out: bool = False
    stderr: str = ""
    stdout: str = ""


class _LocalSofficeSandbox:
    def __init__(self, cwd: Path) -> None:
        self._cwd = cwd

    async def exec_shell(self, cmd: str, *, timeout_s: int) -> _LocalExecResult:
        try:
            res = subprocess.run(
                cmd,
                shell=True,
                cwd=self._cwd,
                timeout=timeout_s,
                check=False,
                capture_output=True,
                text=True,
            )
        except subprocess.TimeoutExpired as exc:
            return _LocalExecResult(
                exit_code=1,
                timed_out=True,
                stderr="" if exc.stderr is None else str(exc.stderr),
            )
        return _LocalExecResult(
            exit_code=res.returncode,
            timed_out=False,
            stderr=res.stderr,
            stdout=res.stdout,
        )


@pytest.mark.integration
@pytest.mark.skipif(shutil.which("soffice") is None, reason="soffice not installed")
@pytest.mark.asyncio
async def test_pdf_export_render_refuses_blank_branded_real_soffice(
    tmp_path: Path,
) -> None:
    pptx_bytes = render_pptx(lower_deck(_blank_branded_checker_deck()))
    pptx_name = "blank-branded.pptx"
    (tmp_path / pptx_name).write_bytes(pptx_bytes)

    ctx = MagicMock()
    ctx.sandbox = _LocalSofficeSandbox(tmp_path)

    ok, err = await convert_to_pdf(ctx, pptx_name)
    assert ok is True, err

    pdf_bytes = (tmp_path / "blank-branded.pdf").read_bytes()
    f = check_export_render(
        "pdf",
        pdf_bytes,
        text=pptx_visible_text(pptx_bytes),
        declared_units=3,
    )
    assert f.non_blank is False and f.ok is False


def test_pptx_export_render_accepts_content_branded_real_render() -> None:
    data = render_pptx(lower_deck(_content_branded_checker_deck()))
    f = check_export_render("pptx", data, declared_units=3)
    assert f.non_blank is True and f.ok is True


# ---- C7 wire: generated image bytes embed into the deck (not [image]) --------


def _tiny_png() -> bytes:
    import io as _io

    from PIL import Image

    buf = _io.BytesIO()
    Image.new("RGB", (32, 24), (123, 200, 90)).save(buf, format="PNG")
    return buf.getvalue()


def test_lower_deck_embeds_image_assets_into_pptx_and_html():
    """C7: lower_deck(image_assets=...) attaches the generated bytes to the image
    element so render_pptx EMBEDS a real picture (ppt/media non-empty) and render_html
    inlines a data-URI — instead of falling back to the [image] placeholder."""
    import zipfile

    from disco.tools.builtin._deck_schema import AuthoredDeck, AuthoredSlide, lower_deck
    from disco.tools.builtin._pptx_render import render_html, render_pptx

    png = _tiny_png()
    authored = AuthoredDeck(
        title="Test",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="full_image", title="Cover", body=[], image_prompt="a green rectangle"
            ),
            AuthoredSlide(type="bullets", title="Plain", body=["no image here"]),
        ],
    )

    # Without assets → no dead "[image]" box. A SIDE slot still gets the themed
    # inline-SVG art fallback (see test_deck_schema's art-fallback test); a
    # FULL-BLEED cover instead drops the art entirely and renders the legible
    # text-only title layout (UI-24) — small white type over a pale generative
    # wash was unreadable.
    deck_plain = lower_deck(authored)
    assert all(el.image_bytes is None for s in deck_plain.slides for el in s.elements)
    plain_html = render_html(deck_plain)
    assert "[image]" not in plain_html
    assert '<h1 class="slide-title"' in plain_html
    assert 'class="slide-image-scrim"' not in plain_html

    # With assets keyed by the image slide's index → embedded.
    deck = lower_deck(authored, image_assets={0: png})
    img_els = [el for s in deck.slides for el in s.elements if el.kind == "image"]
    assert any(el.image_bytes == png for el in img_els), "bytes not attached to image element"

    pptx = render_pptx(deck)
    with zipfile.ZipFile(__import__("io").BytesIO(pptx)) as z:
        media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    assert media, "render_pptx embedded NO media — the image did not make it into the deck"

    html_out = render_html(deck)
    assert "data:image/png;base64," in html_out, "render_html did not inline the image bytes"


def test_lower_deck_image_assets_default_none_is_backcompat():
    """No image_assets (re-theme / editor path) → no image_bytes, current behavior."""
    from disco.tools.builtin._deck_schema import AuthoredDeck, AuthoredSlide, lower_deck

    authored = AuthoredDeck(
        title="T",
        theme="disco-light",
        slides=[AuthoredSlide(type="full_image", title="C", body=[], image_prompt="x")],
    )
    deck = lower_deck(authored)  # no image_assets
    assert all(el.image_bytes is None for s in deck.slides for el in s.elements)


# ---------------------------------------------------------------------------
# W-23 — grouped bullets into one auto-fit frame (no LibreOffice/Google overlap)
# ---------------------------------------------------------------------------


def _long_bullet_deck():
    """A bullets slide whose lines are long enough to wrap under font
    substitution — the case that produced absolute-box overlap before W-23."""
    from disco.tools.builtin._deck_schema import AuthoredDeck, AuthoredSlide, lower_deck

    authored = AuthoredDeck(
        title="W-23 Overlap Repro",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="bullets",
                title="Bullets that wrap",
                body=[
                    "This is a deliberately long first bullet that will wrap onto a "
                    "second line once the brand font is substituted by LibreOffice.",
                    "A second bullet, also long enough to wrap, which used to be "
                    "overlapped by the spillover of the first bullet's second line.",
                    "A third bullet line to make the stacking unmistakable in XML.",
                ],
            )
        ],
    )
    return lower_deck(authored)


def _bullet_frames(slide):
    """Text frames on *slide* whose first paragraph is a bullet ('• ' prefix)."""
    frames = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        paras = shape.text_frame.paragraphs
        if paras and paras[0].runs and paras[0].runs[0].text.startswith("• "):
            frames.append(shape.text_frame)
    return frames


def test_pptx_w23_bullets_grouped_into_one_autofit_frame():
    """W-23: the bullets render as ONE text frame with MULTIPLE bullet paragraphs
    + normAutofit — NOT several overlapping absolute single-bullet textboxes."""
    import io

    from pptx import Presentation

    data = render_pptx(_long_bullet_deck())
    prs = Presentation(io.BytesIO(data))
    slide = prs.slides[0]

    frames = _bullet_frames(slide)
    assert len(frames) == 1, (
        f"expected ONE grouped bullet frame, got {len(frames)} "
        "(per-bullet absolute boxes regressed → overlap risk)"
    )

    tf = frames[0]
    bullet_paras = [p for p in tf.paragraphs if p.runs and p.runs[0].text.startswith("• ")]
    assert len(bullet_paras) == 3, (
        f"expected 3 bullet paragraphs in the one frame, got {len(bullet_paras)}"
    )

    # The frame must carry TEXT_TO_FIT_SHAPE autofit → <a:normAutofit/> in the XML,
    # so a substituted-font wrap shrinks the text instead of spilling.
    xml = tf._txBody.xml
    assert "normAutofit" in xml, "grouped bullet frame is missing <a:normAutofit/> autofit"


def test_pptx_w23_two_column_keeps_separate_column_frames():
    """W-23 grouping is keyed on left/width, so a two_column slide still yields
    one frame per column (bullets differ by `left`)."""
    import io

    from disco.tools.builtin._deck_schema import AuthoredDeck, AuthoredSlide, lower_deck
    from pptx import Presentation

    authored = AuthoredDeck(
        title="Two Col",
        theme="disco-light",
        slides=[
            AuthoredSlide(
                type="two_column",
                title="Split",
                body=["Left one", "Left two", "Right one", "Right two"],
            )
        ],
    )
    data = render_pptx(lower_deck(authored))
    prs = Presentation(io.BytesIO(data))
    frames = _bullet_frames(prs.slides[0])
    assert len(frames) == 2, f"expected 2 column frames, got {len(frames)}"
