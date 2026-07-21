"""W14b — heavy (render-it) validators. poppler (pdftoppm/pdfinfo) is present in CI-lite and
on the VM 201 host, so the PDF rasterizer is unit-tested here; LibreOffice (soffice) is
VM-only, so the pptx renderer's guard path is tested here and the real render runs on VM 201
(proven live against a generated W3-clean deck — see test-record/e2e-full/artifacts/)."""

from __future__ import annotations

import pathlib
import shutil
import struct
import zipfile
import zlib

import pytest
from disco.tools.verify import heavy_validators
from disco.tools.verify.heavy_validators import validate_pdf_renders, validate_pptx_renders

_REPO = pathlib.Path(__file__).resolve().parents[3]
_PDF = _REPO / "packages/tools/tests/fixtures/artifacts/sample_report.pdf"


@pytest.mark.skipif(shutil.which("pdftoppm") is None, reason="poppler not installed")
def test_validate_pdf_renders_real_pdf_nonblank():
    assert validate_pdf_renders(str(_PDF)) == []


@pytest.mark.skipif(shutil.which("pdftoppm") is None, reason="poppler not installed")
def test_validate_pdf_renders_flags_missing():
    problems = validate_pdf_renders("/no/such/file.pdf")
    assert problems and "failed" in problems[0]


def _contentful_pptx_bytes() -> bytes:
    """A real deck with MEANINGFUL extractable content (C9-01: a blank deck renders to a
    text-free PDF, which the PDF validator rightly rejects — the fixture must carry real
    text; the validator must not be weakened)."""
    import io

    from pptx import Presentation

    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Heavy Validator Fixture Deck"
    title.placeholders[1].text = "Rendered by headless LibreOffice in the G17 lane"
    body = prs.slides.add_slide(prs.slide_layouts[1])
    body.shapes.title.text = "Extractable Content"
    body.placeholders[1].text_frame.text = "This bullet proves pdftotext sees real text."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_validate_pptx_renders_guards_when_soffice_absent(
    monkeypatch: pytest.MonkeyPatch, tmp_path: pathlib.Path
) -> None:
    """The unavailable-tool guard, exercised WITHOUT skipping regardless of whether the
    host has LibreOffice (C9-01): an isolated empty PATH makes ``soffice`` unresolvable
    for this test only. The input is a STRUCTURALLY VALID deck so it passes the OPC
    pre-check and actually reaches the soffice-missing guard."""
    empty = tmp_path / "empty-path"
    empty.mkdir()
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(_contentful_pptx_bytes())
    monkeypatch.setenv("PATH", str(empty))
    assert shutil.which("soffice") is None
    problems = validate_pptx_renders(str(deck))
    assert problems == [
        "soffice unavailable — run this heavy validator on the VM 201 evidence host"
    ]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_validate_pptx_renders_real_clean_pptx():
    """When soffice IS present, a CONTENTFUL deck renders with no problems — the render
    path is exercised end-to-end (pinned-filter convert → PDF → page/text checks). Also
    covers the C9-01 verifier defect-2 fix: a RELATIVE ``workdir`` is resolved before the
    profile file-URI is built (it used to raise ``ValueError``)."""
    import os
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        p = pathlib.Path(d) / "deck.pptx"
        p.write_bytes(_contentful_pptx_bytes())
        assert validate_pptx_renders(str(p)) == []
        # A relative workdir must be accepted (resolved), not raise.
        cwd = os.getcwd()
        try:
            os.chdir(d)
            os.mkdir("relwork")
            assert validate_pptx_renders(str(p), workdir="relwork") == []
        finally:
            os.chdir(cwd)


def test_validate_pptx_renders_flags_corrupt_zip():
    """A zip-INVALID .pptx must be flagged BEFORE LibreOffice is even consulted (C9-01):
    the OPC structural pre-check runs first, so ZIP-invalid bytes are rejected on any
    host — no renderer, no skip. (LibreOffice would otherwise EXIT ZERO on "source file
    could not be loaded" and, absent the pinned import filter, recover garbage.)"""
    import tempfile

    with tempfile.TemporaryDirectory() as d:
        p = pathlib.Path(d) / "broken.pptx"
        p.write_bytes(b"this is not a pptx")
        problems = validate_pptx_renders(str(p))
        assert problems, "corrupt pptx produced an empty problem list"
        assert "not a valid zip" in problems[0], problems


def test_validate_pptx_renders_flags_missing_part_that_libreoffice_recovers():
    """C9-01 verifier defect 1: a VALID zip whose slide references a MISSING part
    (a dangling ``_rels`` target) is structurally corrupt even though LibreOffice
    silently recovers and renders it. The OPC integrity pre-check must flag it —
    renderer-independent, so it needs no soffice and never skips."""
    import io
    import tempfile
    import zipfile

    clean = _contentful_pptx_bytes()
    # Drop a slideLayout part while a slide still references it → dangling relationship.
    with zipfile.ZipFile(io.BytesIO(clean)) as zf:
        names = zf.namelist()
        victim = next((n for n in names if "slideLayouts/slideLayout1.xml" in n), None)
        assert victim is not None, "fixture deck has no slideLayout1 to drop"
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for n in names:
                if n == victim:
                    continue
                out.writestr(n, zf.read(n))
    corrupt = buf.getvalue()
    # Sanity: the mutated archive is still a VALID zip (the render path can't rely on
    # a CRC failure here — the reference-integrity check is what must catch it).
    assert zipfile.ZipFile(io.BytesIO(corrupt)).testzip() is None

    with tempfile.TemporaryDirectory() as d:
        p = pathlib.Path(d) / "dangling.pptx"
        p.write_bytes(corrupt)
        problems = validate_pptx_renders(str(p))
        assert problems, "a slide referencing a missing part produced an empty problem list"
        assert any("missing part" in msg for msg in problems), problems


def test_pptx_resource_budget_rejects_before_crc_walk(
    monkeypatch: pytest.MonkeyPatch, tmp_path: pathlib.Path
) -> None:
    """A package outside the declared budget never reaches ``testzip`` (which expands
    every member).  A failed budget check is evidence, not an invitation to keep
    decompressing attacker-controlled bytes."""
    deck = tmp_path / "too-many.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("one.xml", "one")
        zf.writestr("two.xml", "two")
    monkeypatch.setattr(heavy_validators, "_MAX_OPC_MEMBERS", 1)

    def forbidden_crc_walk(self: zipfile.ZipFile) -> str | None:
        raise AssertionError("testzip must not run after a resource-budget rejection")

    monkeypatch.setattr(zipfile.ZipFile, "testzip", forbidden_crc_walk)
    problems = validate_pptx_renders(str(deck))
    assert problems and "too many members" in problems[0]


@pytest.mark.parametrize(
    ("limit_name", "limit", "expected"),
    [
        ("_MAX_OPC_MEMBER_BYTES", 16, "uncompressed-size limit"),
        ("_MAX_OPC_TOTAL_BYTES", 16, "total uncompressed-size limit"),
        ("_MAX_OPC_COMPRESSION_RATIO", 2, "compression-ratio limit"),
    ],
)
def test_pptx_resource_budgets_fail_closed(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: pathlib.Path,
    limit_name: str,
    limit: int,
    expected: str,
) -> None:
    deck = tmp_path / f"{limit_name}.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("payload.xml", "A" * 1_024)
    monkeypatch.setattr(heavy_validators, limit_name, limit)
    problems = validate_pptx_renders(str(deck))
    assert problems and expected in problems[0]


def test_oversized_relationship_xml_is_never_parsed(
    monkeypatch: pytest.MonkeyPatch, tmp_path: pathlib.Path
) -> None:
    deck = tmp_path / "large-rels.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("ppt/_rels/presentation.xml.rels", "<Relationships>" + " " * 128)
    monkeypatch.setattr(heavy_validators, "_MAX_OPC_RELATIONSHIP_BYTES", 32)

    def forbidden_parse(_data: bytes) -> object:
        raise AssertionError("oversized relationship XML must not be parsed")

    monkeypatch.setattr(heavy_validators.ET, "fromstring", forbidden_parse)
    problems = validate_pptx_renders(str(deck))
    assert problems and "XML size limit" in problems[0]


def test_relationship_dtd_entity_is_rejected_without_expansion(tmp_path: pathlib.Path) -> None:
    deck = tmp_path / "entity.pptx"
    rels = b"""<?xml version="1.0"?>
<!DOCTYPE Relationships [
<!ENTITY a "1234567890">
<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">
]>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="x" Target="&b;"/>
</Relationships>
"""
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("ppt/_rels/presentation.xml.rels", rels)
    problems = heavy_validators._opc_integrity_problems(str(deck))
    assert problems == [
        "pptx relationship part ppt/_rels/presentation.xml.rels contains a forbidden DTD/entity"
    ]
    assert len(problems[0]) < 200


def test_forged_declared_size_cannot_hide_deflate_expansion(tmp_path: pathlib.Path) -> None:
    """Patch both local and central metadata to claim one output byte with its valid
    prefix CRC.  Python's ordinary ZipExtFile/testzip trusts that declaration; the raw
    bounded stream proof must still catch the hidden deflate output."""
    deck = tmp_path / "forged-size.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("payload.bin", b"A" * (2 * 1024 * 1024))
    raw = bytearray(deck.read_bytes())
    local = raw.index(b"PK\x03\x04")
    central = raw.index(b"PK\x01\x02")
    prefix_crc = zlib.crc32(b"A") & 0xFFFFFFFF
    struct.pack_into("<L", raw, local + 14, prefix_crc)
    struct.pack_into("<L", raw, local + 22, 1)
    struct.pack_into("<L", raw, central + 16, prefix_crc)
    struct.pack_into("<L", raw, central + 24, 1)
    deck.write_bytes(raw)
    with zipfile.ZipFile(deck) as zf:
        assert zf.testzip() is None
        assert zf.read("payload.bin") == b"A"
    problems = heavy_validators._opc_integrity_problems(str(deck))
    assert problems and "expands beyond its declared size" in problems[0]


def test_duplicate_opc_part_names_are_rejected(tmp_path: pathlib.Path) -> None:
    deck = tmp_path / "duplicate.pptx"
    with pytest.warns(UserWarning, match="Duplicate name"):
        with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("ppt/_rels/presentation.xml.rels", b"<Relationships/>")
            zf.writestr("ppt/_rels/presentation.xml.rels", b"<Relationships/>")
    assert heavy_validators._opc_integrity_problems(str(deck)) == [
        "pptx archive contains duplicate part names"
    ]


def test_utf16_relationship_dtd_is_rejected_and_comment_text_is_allowed(
    tmp_path: pathlib.Path,
) -> None:
    malicious = tmp_path / "utf16-dtd.pptx"
    dtd_xml = """<?xml version="1.0" encoding="utf-16"?>
<!DOCTYPE Relationships [<!ENTITY x "expanded">]>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="r1" Type="x" Target="&x;" />
</Relationships>""".encode("utf-16")
    with zipfile.ZipFile(malicious, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("ppt/_rels/presentation.xml.rels", dtd_xml)
    assert "forbidden DTD/entity" in heavy_validators._opc_integrity_problems(str(malicious))[0]

    harmless = tmp_path / "comment.pptx"
    comment_xml = b"""<Relationships
xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<!-- documentation spells <!DOCTYPE but does not declare one -->
</Relationships>"""
    with zipfile.ZipFile(harmless, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("ppt/_rels/presentation.xml.rels", comment_xml)
    assert heavy_validators._opc_integrity_problems(str(harmless)) == []


def test_malformed_deflate_and_local_flag_disagreement_fail_as_data(
    tmp_path: pathlib.Path,
) -> None:
    malformed = tmp_path / "malformed-deflate.pptx"
    with zipfile.ZipFile(malformed, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("payload.bin", b"payload" * 1_000)
    raw = bytearray(malformed.read_bytes())
    local = raw.index(b"PK\x03\x04")
    name_size = struct.unpack_from("<H", raw, local + 26)[0]
    extra_size = struct.unpack_from("<H", raw, local + 28)[0]
    data_offset = local + 30 + name_size + extra_size
    raw[data_offset] ^= 0xFF
    malformed.write_bytes(raw)
    problems = heavy_validators._opc_integrity_problems(str(malformed))
    assert problems and "safely inspected" in problems[0]

    flags = tmp_path / "flag-mismatch.pptx"
    with zipfile.ZipFile(flags, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("payload.bin", b"payload")
    raw = bytearray(flags.read_bytes())
    local = raw.index(b"PK\x03\x04")
    local_flags = struct.unpack_from("<H", raw, local + 6)[0]
    struct.pack_into("<H", raw, local + 6, local_flags ^ 0x0800)
    flags.write_bytes(raw)
    assert heavy_validators._opc_integrity_problems(str(flags)) == [
        "pptx member payload.bin has inconsistent local metadata"
    ]


def test_relationship_diagnostics_are_globally_bounded(tmp_path: pathlib.Path) -> None:
    deck = tmp_path / "many-invalid-rels.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        for index in range(250):
            zf.writestr(f"ppt/parts/{index}/_rels/item.xml.rels", b"<not-xml")
    problems = heavy_validators._opc_integrity_problems(str(deck))
    assert len(problems) == heavy_validators._MAX_OPC_PROBLEMS
    assert problems[-1] == "pptx has additional structural problems"


def test_forced_zip64_local_sizes_are_validated_without_false_rejection(
    tmp_path: pathlib.Path,
) -> None:
    deck = tmp_path / "forced-zip64.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED, allowZip64=True) as zf:
        with zf.open("payload.xml", "w", force_zip64=True) as member:
            member.write(b"<payload>valid</payload>")
    assert heavy_validators._opc_integrity_problems(str(deck)) == []


def test_local_and_central_member_names_must_agree(tmp_path: pathlib.Path) -> None:
    deck = tmp_path / "name-mismatch.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("payload.bin", b"payload")
    raw = bytearray(deck.read_bytes())
    local = raw.index(b"PK\x03\x04")
    name_offset = local + 30
    assert raw[name_offset : name_offset + len(b"payload.bin")] == b"payload.bin"
    raw[name_offset] = ord("q")
    deck.write_bytes(raw)
    assert heavy_validators._opc_integrity_problems(str(deck)) == [
        "pptx member payload.bin disagrees on its local filename"
    ]


def test_unicode_member_name_matches_under_zip_utf8_flag(tmp_path: pathlib.Path) -> None:
    deck = tmp_path / "unicode-name.pptx"
    with zipfile.ZipFile(deck, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("ppt/média.xml", b"<payload />")
    assert heavy_validators._opc_integrity_problems(str(deck)) == []
