"""Raw python-pptx primitive drawing helpers — moved verbatim out of ``_pptx_render.py``.

Low-level shape/text/color helpers (``_add_textbox``, ``_set_run_style``,
``_add_accent_bar``, ...), the default-template brand chrome (wordmark +
cover colophon), the generative-art placeholder for an empty image slot, and
the themed-text convenience wrapper used by the C1 native-pptx archetype
renderers. Every python-pptx-typed function keeps its pre-existing
type-checker suppression comment exactly as it was; none was added here.
"""

from __future__ import annotations

from pathlib import Path

from disco.core.brand.tokens import Theme

from .._deck_schema import Slide
from ._geometry import _BRAND_DOT, _CW, _MARGIN, _SLIDE_H, _SLIDE_W


def _rgb_from_hex(hex_color: str):  # type: ignore[return]
    """Return a pptx RGBColor from a ``#rrggbb`` string."""
    from pptx.dml.color import RGBColor  # lazy import — avoids load-time dep

    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _first_font(stack: str) -> str:
    """Extract the first font name from a CSS font-family stack.

    ``"'Fraunces',Georgia,serif"`` → ``"Fraunces"``.
    """
    return stack.split(",")[0].strip().strip("'\"")


def _font_dir() -> Path:
    """Absolute path to the bundled OFL TTF fonts (via disco.core.brand)."""
    import importlib.resources

    pkg = importlib.resources.files("disco.core.brand")
    return Path(str(pkg / "fonts"))


def _set_slide_bg(slide, hex_color: str) -> None:  # type: ignore[type-arg]
    """Fill a python-pptx slide background with a solid hex color."""
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = _rgb_from_hex(hex_color)


def _add_textbox(slide, left: int, top: int, w: int, h: int):  # type: ignore[return,type-arg]
    """Add a textbox at absolute EMU coordinates and return its text_frame.

    W-23: every frame gets ``TEXT_TO_FIT_SHAPE`` autofit (emits ``<a:normAutofit/>``
    in the slide XML) so that when a viewer substitutes the brand font (not
    installed in LibreOffice / Google Slides → different metrics → an extra wrap
    line) the text SHRINKS to stay inside the box instead of spilling out and
    overlapping the next element. Applies the safety to ordinary title / subtitle
    / kicker frames; grouped bullet frames add margins + anchor on top of this.
    """
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu

    shape = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(w), Emu(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return tf


def _set_run_style(
    run,  # type: ignore[type-arg]
    text: str,
    font_name: str,
    size_pt: float,
    hex_color: str,
    bold: bool = False,
    italic: bool = False,
) -> None:
    """Apply text + font styling to a python-pptx run."""
    from pptx.util import Pt

    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = _rgb_from_hex(hex_color)
    run.font.bold = bold
    run.font.italic = italic


def _add_accent_bar(slide, left: int, top: int, w: int, hex_color: str) -> None:
    """Add a thin horizontal accent rule (0.1 in tall)."""
    from pptx.util import Emu

    # Use add_shape with RECTANGLE (auto-shape type 1)
    bar = slide.shapes.add_shape(1, Emu(left), Emu(top), Emu(w), Emu(91_440))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb_from_hex(hex_color)
    bar.line.fill.background()  # no border line


# ---------------------------------------------------------------------------
# Brand marks — the default Disco template chrome that mirrors the PDF report:
#   • a small "Disco." wordmark on every slide (the PDF's running header), and
#   • the "disco — Latin · verb" definition-mark colophon on the cover/title
#     slide (the PDF cover's signature).
# Drawn at RENDER time (not as deck Elements) so the HTML title/bullet
# heuristics in _html_for_c1_slide stay untouched. Gated on theme.branded, so
# the neutral theme renders clean — exactly like serialize_pdf's brand gate.
# ---------------------------------------------------------------------------


def _draw_wordmark(prs_slide, theme: Theme) -> None:  # type: ignore[type-arg]
    """Small 'Disco' + accent dot in the top margin band — the PDF running wordmark.

    Kept ENTIRELY ABOVE the title strip (titles start at _MARGIN = 0.5in): the box
    bottom must stay < _MARGIN so a long, right-reaching title never collides with it."""
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu

    dot = _BRAND_DOT
    box_w = 1_500_000
    box_h = 259_080  # 0.283in
    top = 91_440  # 0.1in  → bottom 0.383in, clear of the 0.5in title line
    right = _SLIDE_W - _MARGIN
    text_box_right = right - dot - 36_576  # leave room for the trailing dot
    tf = _add_textbox(prs_slide, text_box_right - box_w, top, box_w, box_h)
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT
    run = para.add_run()
    _set_run_style(run, "Disco", _first_font(theme.font_display), 10.0, theme.text_faint)
    # Accent dot, aligned with the small wordmark text baseline.
    dot_top = top + 60_000
    d = prs_slide.shapes.add_shape(1, Emu(right - dot), Emu(dot_top), Emu(dot), Emu(dot))
    d.fill.solid()
    d.fill.fore_color.rgb = _rgb_from_hex(theme.accent)
    d.line.fill.background()


def _draw_cover_colophon(prs_slide, theme: Theme) -> None:  # type: ignore[type-arg]
    """The 'disco — Latin · verb' definition mark at the bottom of the cover/title
    slide — the PDF cover's signature colophon (headword + POS + gloss + etymology)."""
    from pptx.enum.text import PP_ALIGN

    top = int(_SLIDE_H * 0.74)
    # Short accent rule above the mark.
    _add_accent_bar(prs_slide, _MARGIN, top, 640_080, theme.accent)

    # Headword "disco" + part-of-speech / IPA on one line (two runs, two styles).
    hw_top = top + 140_000
    tf = _add_textbox(prs_slide, _MARGIN, hw_top, _CW, 360_000)
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.LEFT
    r_hw = para.add_run()
    _set_run_style(r_hw, "disco   ", _first_font(theme.font_display), 16.0, theme.text)
    r_pos = para.add_run()
    _set_run_style(
        r_pos,
        "LATIN · VERB    /ˈdɪs.koː/",
        _first_font(theme.font_ui),
        8.5,
        theme.text_faint,
    )

    # Italic gloss.
    gloss_top = hw_top + 430_000
    tf2 = _add_textbox(prs_slide, _MARGIN, gloss_top, _CW, 330_000)
    r_gloss = tf2.paragraphs[0].add_run()
    _set_run_style(
        r_gloss,
        "“I learn; I become acquainted with.”",
        _first_font(theme.font_reading),
        12.0,
        theme.text_muted,
        italic=True,
    )

    # Etymology root.
    root_top = gloss_top + 365_000
    tf3 = _add_textbox(prs_slide, _MARGIN, root_top, _CW, 300_000)
    r_root = tf3.paragraphs[0].add_run()
    _set_run_style(
        r_root,
        "from discere — to learn",
        _first_font(theme.font_ui),
        8.5,
        theme.text_faint,
    )


def _draw_brand_marks(prs_slide, slide: Slide, theme: Theme) -> None:  # type: ignore[type-arg]
    """Apply the default-template brand chrome to a rendered slide (no-op when the
    theme is unbranded). The colophon rides only the cover/title slide."""
    if not theme.branded:
        return
    _draw_wordmark(prs_slide, theme)
    # Gate the colophon on the RESOLVED layout (not the authored type) so only a
    # true rendered title cover gets it — e.g. type="title"+layout_hint="closing"
    # must NOT, and the cover slide always does.
    if slide.layout == "title":
        _draw_cover_colophon(prs_slide, theme)


def _image_placeholder(prs_slide, left: int, top: int, w: int, h: int, theme: Theme) -> None:  # type: ignore[type-arg]
    """Themed generative-art fallback for an empty image slot (image gen failed
    or is unconfigured). Replaces the old flat gray "[image]" box (Dylan
    2026-07-07: 'if image gen doesn't work then generate SVGs') — an abstract,
    deterministic geometric composition in the deck's own palette, drawn with
    native vector shapes (the PPTX analogue of the HTML renderer's inline-SVG
    fallback). Variant is seeded from the slot geometry so repeated slots in one
    deck differ but re-renders are byte-stable."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu

    # Base field.
    base = prs_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(left), Emu(top), Emu(w), Emu(h))
    base.fill.solid()
    base.fill.fore_color.rgb = _rgb_from_hex(theme.surface_1)
    base.line.color.rgb = _rgb_from_hex(theme.hairline)

    # Deterministic per-slot variant; distinct primes mix the geometry so a
    # cover slot and a side slot land on different compositions.
    variant = (left * 7 + top * 13 + w * 3 + h * 5) // 9_525 % 3
    r_big = min(w, h) * 2 // 3
    r_mid = r_big // 2
    r_dot = max(r_big // 7, 91_440)

    def _circle(x: int, y: int, r: int, hex_color: str, outline: bool = False) -> None:
        c = prs_slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Emu(x - r // 2), Emu(y - r // 2), Emu(r), Emu(r)
        )
        if outline:
            c.fill.background()
            c.line.color.rgb = _rgb_from_hex(hex_color)
            c.line.width = Emu(19_050)
        else:
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb_from_hex(hex_color)
            c.line.fill.background()

    def _bar(x: int, y: int, bw: int, bh: int, hex_color: str) -> None:
        b = prs_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(bw), Emu(bh))
        b.fill.solid()
        b.fill.fore_color.rgb = _rgb_from_hex(hex_color)
        b.line.fill.background()

    if variant == 0:
        # Off-center large disc + outlined echo + accent dot.
        _circle(left + w * 2 // 3, top + h // 3, r_big, theme.surface_2)
        _circle(
            left + w * 2 // 3, top + h // 3, r_big + r_mid // 2, theme.hairline_strong, outline=True
        )
        _circle(left + w // 4, top + h * 3 // 4, r_dot, theme.accent)
        _bar(left + w // 8, top + h * 5 // 6, w // 3, 27_432, theme.hairline_strong)
    elif variant == 1:
        # Horizon band + rising disc (accent) + faint moon.
        _bar(left, top + h * 2 // 3, w, h // 24 + 18_288, theme.surface_2)
        _circle(left + w // 3, top + h * 2 // 3, r_mid, theme.accent)
        _circle(left + w * 4 // 5, top + h // 4, r_dot * 2, theme.surface_2)
    else:
        # Column rhythm + accent square.
        for i in range(4):
            _bar(
                left + w * (2 * i + 1) // 9,
                top + h // 5,
                w // 18 + 9_525,
                h * 3 // 5,
                theme.surface_2,
            )
        sq = max(r_dot, 137_160)
        _bar(left + w * 3 // 4, top + h * 2 // 3, sq, sq, theme.accent)


def _add_pptx_text(
    prs_slide,
    text: str,
    *,
    left: int,
    top: int,
    width: int,
    height: int,
    theme: Theme,
    font_name: str,
    size_pt: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
    align: str = "LEFT",
) -> None:
    from pptx.enum.text import PP_ALIGN

    align_map = {"LEFT": PP_ALIGN.LEFT, "CENTER": PP_ALIGN.CENTER, "RIGHT": PP_ALIGN.RIGHT}
    tf = _add_textbox(prs_slide, left, top, width, height)
    p = tf.paragraphs[0]
    p.alignment = align_map.get(align, PP_ALIGN.LEFT)
    run = p.add_run()
    _set_run_style(run, text, font_name, size_pt, color, bold=bold, italic=italic)
