"""C1 native-pptx slide-archetype renderers (big_number / quote / timeline /
two_by_two) and their dispatcher.

Moved verbatim out of ``_pptx_render.py``.
"""

from __future__ import annotations

from disco.core.brand.tokens import Theme

from .._deck_schema import Slide
from ._geometry import _CW, _MARGIN, _SLIDE_W
from ._pptx_primitives import _add_pptx_text, _first_font, _rgb_from_hex
from ._slide_archetype_shared import _body_strings, _slide_archetype, _split_timeline_label


def _render_big_number_pptx(prs_slide, slide: Slide, theme: Theme) -> None:  # type: ignore[type-arg]
    body = _body_strings(slide)
    figure = body[0] if body else slide.title
    support = body[1] if len(body) > 1 else (slide.title if body else "")
    if slide.title and slide.title != support:
        _add_pptx_text(
            prs_slide,
            slide.title,
            left=_MARGIN,
            top=_MARGIN,
            width=_CW,
            height=365_760,
            theme=theme,
            font_name=_first_font(theme.font_ui),
            size_pt=18.0,
            color=theme.text_muted,
            bold=True,
        )
    _add_pptx_text(
        prs_slide,
        figure,
        left=_MARGIN,
        top=1_371_600,
        width=_CW,
        height=2_286_000,
        theme=theme,
        font_name=_first_font(theme.font_display),
        size_pt=150.0,
        color=theme.accent,
        bold=True,
        align="CENTER",
    )
    if support:
        _add_pptx_text(
            prs_slide,
            support,
            left=int(_MARGIN * 1.5),
            top=3_657_600,
            width=_SLIDE_W - int(_MARGIN * 3),
            height=731_520,
            theme=theme,
            font_name=_first_font(theme.font_reading),
            size_pt=24.0,
            color=theme.text,
            align="CENTER",
        )


def _render_quote_pptx(prs_slide, slide: Slide, theme: Theme) -> None:  # type: ignore[type-arg]
    body = _body_strings(slide)
    quote = body[0] if body else slide.title
    attribution = body[1] if len(body) > 1 else ""
    _add_pptx_text(
        prs_slide,
        "“",
        left=_MARGIN,
        top=914_400,
        width=914_400,
        height=1_371_600,
        theme=theme,
        font_name=_first_font(theme.font_display),
        size_pt=120.0,
        color=theme.accent,
        bold=True,
    )
    _add_pptx_text(
        prs_slide,
        quote,
        left=1_371_600,
        top=1_280_160,
        width=_SLIDE_W - 2_286_000,
        height=2_743_200,
        theme=theme,
        font_name=_first_font(theme.font_reading),
        size_pt=34.0,
        color=theme.text,
        italic=True,
    )
    if attribution:
        _add_pptx_text(
            prs_slide,
            attribution,
            left=1_371_600,
            top=4_206_240,
            width=_SLIDE_W - 2_286_000,
            height=457_200,
            theme=theme,
            font_name=_first_font(theme.font_ui),
            size_pt=16.0,
            color=theme.text_muted,
            bold=True,
        )


def _render_timeline_pptx(prs_slide, slide: Slide, theme: Theme) -> None:  # type: ignore[type-arg]
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu

    items = _body_strings(slide) or [slide.title]
    title = slide.title
    if title:
        _add_pptx_text(
            prs_slide,
            title,
            left=_MARGIN,
            top=_MARGIN,
            width=_CW,
            height=548_640,
            theme=theme,
            font_name=_first_font(theme.font_ui),
            size_pt=24.0,
            color=theme.text,
            bold=True,
        )
    margin = 914_400
    y = 3_337_560
    spine = prs_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(margin), Emu(y), Emu(_SLIDE_W - 2 * margin), Emu(18_288)
    )
    spine.fill.solid()
    spine.fill.fore_color.rgb = _rgb_from_hex(theme.accent)
    spine.line.fill.background()
    denom = max(1, len(items) - 1)
    for i, item in enumerate(items):
        x = margin + int(i * ((_SLIDE_W - 2 * margin) / denom))
        node = prs_slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Emu(x - 54_864), Emu(y - 54_864), Emu(109_728), Emu(109_728)
        )
        node.fill.solid()
        node.fill.fore_color.rgb = _rgb_from_hex(theme.accent)
        node.line.fill.background()
        head, detail = _split_timeline_label(item)
        label = head if not detail else f"{head}\n{detail}"
        label_top = y - 914_400 if i % 2 == 0 else y + 228_600
        _add_pptx_text(
            prs_slide,
            label,
            left=x - 914_400,
            top=label_top,
            width=1_828_800,
            height=731_520,
            theme=theme,
            font_name=_first_font(theme.font_reading),
            size_pt=14.0,
            color=theme.text,
            align="CENTER",
        )


def _render_two_by_two_pptx(prs_slide, slide: Slide, theme: Theme) -> None:  # type: ignore[type-arg]
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu

    body = _body_strings(slide)
    # Axis labels ONLY when the model authored them (6+ body lines: x, y, 4 cells).
    # The old fallback INVENTED "Higher impact/certainty" on 4-cell content — a
    # fabricated analytical framing the quadrants never plotted (no false
    # affordances; gauntlet e-news visual review 2026-07-07).
    axis_x = body[0] if len(body) >= 6 else None
    axis_y = body[1] if len(body) >= 6 else None
    cells = body[2:6] if len(body) >= 6 else body[:4]
    while len(cells) < 4:
        cells.append("")
    _add_pptx_text(
        prs_slide,
        slide.title,
        left=_MARGIN,
        top=_MARGIN,
        width=_CW,
        height=548_640,
        theme=theme,
        font_name=_first_font(theme.font_ui),
        size_pt=24.0,
        color=theme.text,
        bold=True,
    )
    left = 1_371_600
    top = 1_371_600
    grid_w = _SLIDE_W - 2_743_200
    grid_h = 4_389_120
    cell_w = grid_w // 2
    cell_h = grid_h // 2
    for idx, label in enumerate(cells):
        col = idx % 2
        row = idx // 2
        box = prs_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Emu(left + col * cell_w),
            Emu(top + row * cell_h),
            Emu(cell_w),
            Emu(cell_h),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = _rgb_from_hex(theme.surface_1)
        box.line.color.rgb = _rgb_from_hex(theme.hairline)
        _add_pptx_text(
            prs_slide,
            label,
            left=left + col * cell_w + 137_160,
            top=top + row * cell_h + 137_160,
            width=cell_w - 274_320,
            height=cell_h - 274_320,
            theme=theme,
            font_name=_first_font(theme.font_reading),
            size_pt=18.0,
            color=theme.text,
            bold=True,
        )
    if axis_x:
        _add_pptx_text(
            prs_slide,
            axis_x,
            left=left + grid_w - 1_828_800,
            top=top + grid_h + 91_440,
            width=1_828_800,
            height=274_320,
            theme=theme,
            font_name=_first_font(theme.font_ui),
            size_pt=10.0,
            color=theme.accent,
            bold=True,
            align="RIGHT",
        )
    if axis_y:
        _add_pptx_text(
            prs_slide,
            axis_y,
            left=left - 914_400,
            top=top,
            width=822_960,
            height=274_320,
            theme=theme,
            font_name=_first_font(theme.font_ui),
            size_pt=10.0,
            color=theme.accent,
            bold=True,
            align="RIGHT",
        )


def _render_archetype_pptx(prs_slide, slide: Slide, theme: Theme) -> bool:  # type: ignore[type-arg]
    archetype = _slide_archetype(slide)
    if archetype == "big_number":
        _render_big_number_pptx(prs_slide, slide, theme)
        return True
    if archetype == "quote":
        _render_quote_pptx(prs_slide, slide, theme)
        return True
    if archetype == "timeline":
        _render_timeline_pptx(prs_slide, slide, theme)
        return True
    if archetype == "two_by_two":
        _render_two_by_two_pptx(prs_slide, slide, theme)
        return True
    return False
