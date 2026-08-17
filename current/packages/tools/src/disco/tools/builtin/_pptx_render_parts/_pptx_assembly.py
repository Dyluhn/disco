"""render_pptx — the per-deck native-pptx assembly orchestrator.

Moved verbatim out of ``_pptx_render.py``. ``render_pptx`` (the public
compat/export facade, unmoved) converts a ``MinimalDeck`` to a C1 ``Deck``
then calls ``_render_pptx_c1`` here.
"""

from __future__ import annotations

import io

from .._deck_schema import Deck
from ._geometry import _SLIDE_H, _SLIDE_W
from ._pptx_archetypes import _render_archetype_pptx
from ._pptx_c1_elements import _render_slide_elements
from ._pptx_primitives import _draw_brand_marks, _set_slide_bg


def _render_pptx_c1(deck: Deck) -> bytes:
    """Render a C1 Deck to .pptx bytes."""
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(_SLIDE_W)
    prs.slide_height = Emu(_SLIDE_H)
    blank_layout = prs.slide_layouts[6]

    for slide in deck.slides:
        prs_slide = prs.slides.add_slide(blank_layout)
        _set_slide_bg(prs_slide, deck.theme.bg)

        # C8: chart/table slides delegate to native chart/table shapes
        if slide.chart is not None:
            from disco.tools.builtin._c8_chart_layouts import layout_chart_slide_pptx

            layout_chart_slide_pptx(prs_slide, slide, deck.theme)
        elif slide.table is not None:
            from disco.tools.builtin._c8_chart_layouts import layout_table_slide_pptx

            layout_table_slide_pptx(prs_slide, slide, deck.theme)
        else:
            if not _render_archetype_pptx(prs_slide, slide, deck.theme):
                _render_slide_elements(prs_slide, slide, deck.theme)

            if slide.notes:
                _ntf = prs_slide.notes_slide.notes_text_frame
                if _ntf is not None:
                    _ntf.text = slide.notes

        # Default-template brand chrome (wordmark on every slide; colophon on the
        # cover) — applied after content so it overlays cleanly. Chart/table slides
        # get the wordmark too (theme-gated inside).
        _draw_brand_marks(prs_slide, slide, deck.theme)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
