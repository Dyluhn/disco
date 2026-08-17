"""MinimalDeck (compat-shim) per-layout native-pptx renderers.

Moved verbatim out of ``_pptx_render.py``. ``render_pptx`` always converts a
``MinimalDeck`` to a C1 ``Deck`` first (via ``_minimal_to_authored`` →
``lower_deck``) before rendering, so these per-layout functions and
``_MINIMAL_LAYOUT_FNS`` are not currently wired into any call path — kept as a
pre-existing (unused) compatibility surface, unchanged by this split.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from disco.core.brand.tokens import Theme

from ._geometry import _CH, _CW, _MARGIN, _SLIDE_H, _TITLE_H
from ._pptx_primitives import (
    _add_accent_bar,
    _add_textbox,
    _first_font,
    _image_placeholder,
    _set_run_style,
    _set_slide_bg,
)

if TYPE_CHECKING:
    from .._pptx_render import DeckSlide


def _layout_title_slide(prs_slide, slide: DeckSlide, theme: Theme) -> None:  # type: ignore[type-arg]
    """Title-slide layout: large centered display title + optional subtitle.

    Visual structure:
      [vertical center]
        accent bar (2.5 in wide)
        Title      (Fraunces, 52 pt)
        Subtitle   (Newsreader italic, 24 pt, first bullet if present)

    All elements absolutely positioned in EMU on the 16:9 canvas.
    """
    from pptx.enum.text import PP_ALIGN

    _set_slide_bg(prs_slide, theme.bg)

    # Vertical origin (centered block starting at ~30% height)
    v_origin = int(_SLIDE_H * 0.28)

    # Accent bar
    bar_w = 2_286_000  # 2.5 in
    _add_accent_bar(prs_slide, _MARGIN, v_origin, bar_w, theme.accent)

    # Title textbox
    title_top = v_origin + 182_880  # 0.2 in below bar
    title_h = 1_828_800  # 2 in
    tf_title = _add_textbox(prs_slide, _MARGIN, title_top, _CW, title_h)
    p = tf_title.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    _set_run_style(
        run,
        slide.title,
        _first_font(theme.font_display),
        52,
        theme.text,
        bold=False,
    )

    # Subtitle (first bullet, if any)
    if slide.bullets:
        sub_top = title_top + title_h
        sub_h = 914_400  # 1 in
        tf_sub = _add_textbox(prs_slide, _MARGIN, sub_top, _CW, sub_h)
        p_sub = tf_sub.paragraphs[0]
        p_sub.alignment = PP_ALIGN.LEFT
        run_sub = p_sub.add_run()
        _set_run_style(
            run_sub,
            slide.bullets[0],
            _first_font(theme.font_reading),
            24,
            theme.text_muted,
            italic=True,
        )

    # Speaker notes
    if slide.notes:
        prs_slide.notes_slide.notes_text_frame.text = slide.notes


def _layout_bullets_slide(prs_slide, slide: DeckSlide, theme: Theme) -> None:  # type: ignore[type-arg]
    """Standard title + bullets layout.

    Visual structure:
      Title (Schibsted Grotesk, 32 pt) — top 1 in, left-flush
      Accent underline bar (full content width, 2 pt thick)
      Bullets (Newsreader, 20 pt) — below bar, indented
    """
    from pptx.enum.text import PP_ALIGN

    _set_slide_bg(prs_slide, theme.bg)

    # Title
    tf_title = _add_textbox(prs_slide, _MARGIN, _MARGIN, _CW, _TITLE_H)
    p_title = tf_title.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    run_title = p_title.add_run()
    _set_run_style(
        run_title,
        slide.title,
        _first_font(theme.font_ui),
        32,
        theme.text,
        bold=True,
    )

    # Accent bar under title (thin, full width)
    bar_top = _MARGIN + _TITLE_H + 45_720  # 0.05 in gap
    _add_accent_bar(prs_slide, _MARGIN, bar_top, _CW, theme.accent)

    # Bullets textbox — occupies remaining vertical space
    bullet_top = bar_top + 91_440 + 91_440  # 0.1 in bar + 0.1 in gap
    bullet_h = _SLIDE_H - bullet_top - _MARGIN
    tf_bullets = _add_textbox(prs_slide, _MARGIN + 228_600, bullet_top, _CW - 228_600, bullet_h)
    tf_bullets.word_wrap = True

    for i, bullet in enumerate(slide.bullets):
        if i == 0:
            p = tf_bullets.paragraphs[0]
        else:
            p = tf_bullets.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        # Bullet marker as a run prefix
        run = p.add_run()
        _set_run_style(
            run,
            f"• {bullet}",
            _first_font(theme.font_reading),
            20,
            theme.text,
        )

    if slide.notes:
        prs_slide.notes_slide.notes_text_frame.text = slide.notes


def _layout_section_slide(prs_slide, slide: DeckSlide, theme: Theme) -> None:  # type: ignore[type-arg]
    """Section-divider layout: vertically centered section title + optional subtitle.

    Renders with surface_1 background to visually separate sections.
    """
    from pptx.enum.text import PP_ALIGN

    _set_slide_bg(prs_slide, theme.surface_1)

    v_center = _SLIDE_H // 2

    # Accent bar (left-aligned, 1.5 in wide)
    bar_w = 1_371_600  # 1.5 in
    bar_top = v_center - 228_600  # 0.25 in above center
    _add_accent_bar(prs_slide, _MARGIN, bar_top, bar_w, theme.accent)

    # Section label (small allcaps kicker above the title)
    kicker_top = bar_top - 457_200  # 0.5 in above bar
    tf_kicker = _add_textbox(prs_slide, _MARGIN, kicker_top, _CW, 457_200)
    p_kicker = tf_kicker.paragraphs[0]
    p_kicker.alignment = PP_ALIGN.LEFT
    run_kicker = p_kicker.add_run()
    _set_run_style(
        run_kicker,
        "SECTION",
        _first_font(theme.font_ui),
        9,
        theme.accent,
        bold=True,
    )

    # Title (Fraunces, 40 pt)
    title_top = bar_top + 182_880  # 0.2 in below bar
    title_h = 1_371_600  # 1.5 in
    tf_title = _add_textbox(prs_slide, _MARGIN, title_top, _CW, title_h)
    p_title = tf_title.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    run_title = p_title.add_run()
    _set_run_style(
        run_title,
        slide.title,
        _first_font(theme.font_display),
        40,
        theme.text,
    )

    # Optional subtitle (first bullet)
    if slide.bullets:
        sub_top = title_top + title_h
        tf_sub = _add_textbox(prs_slide, _MARGIN, sub_top, _CW, 914_400)
        p_sub = tf_sub.paragraphs[0]
        p_sub.alignment = PP_ALIGN.LEFT
        run_sub = p_sub.add_run()
        _set_run_style(
            run_sub,
            slide.bullets[0],
            _first_font(theme.font_reading),
            18,
            theme.text_muted,
            italic=True,
        )

    if slide.notes:
        prs_slide.notes_slide.notes_text_frame.text = slide.notes


def _layout_image_right_slide(prs_slide, slide: DeckSlide, theme: Theme) -> None:  # type: ignore[type-arg]
    """Title + bullets on left half, image placeholder on right half.

    When image_url is None (C7 image-gen not yet wired), renders a styled
    placeholder box so the layout is correct and C7 can fill it in later.
    """
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu

    _set_slide_bg(prs_slide, theme.bg)

    half_w = _CW // 2 - 91_440  # slight gap between halves

    # Left half: title + bullets
    tf_title = _add_textbox(prs_slide, _MARGIN, _MARGIN, half_w, _TITLE_H)
    p_title = tf_title.paragraphs[0]
    p_title.alignment = PP_ALIGN.LEFT
    run_title = p_title.add_run()
    _set_run_style(
        run_title,
        slide.title,
        _first_font(theme.font_ui),
        28,
        theme.text,
        bold=True,
    )

    bar_top = _MARGIN + _TITLE_H + 45_720
    _add_accent_bar(prs_slide, _MARGIN, bar_top, half_w, theme.accent)

    bullet_top = bar_top + 182_880
    bullet_h = _SLIDE_H - bullet_top - _MARGIN
    tf_b = _add_textbox(prs_slide, _MARGIN + 228_600, bullet_top, half_w - 228_600, bullet_h)
    tf_b.word_wrap = True
    for i, bullet in enumerate(slide.bullets):
        p = tf_b.paragraphs[0] if i == 0 else tf_b.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        _set_run_style(run, f"• {bullet}", _first_font(theme.font_reading), 18, theme.text)

    # Right half: image or placeholder
    img_left = _MARGIN + half_w + 182_880
    img_w = _CW - half_w - 182_880
    img_h = _CH

    if slide.image_url:
        try:
            prs_slide.shapes.add_picture(
                slide.image_url,
                Emu(img_left),
                Emu(_MARGIN),
                Emu(img_w),
                Emu(img_h),
            )
        except Exception:
            # Fall back to placeholder if image can't be loaded
            _image_placeholder(prs_slide, img_left, _MARGIN, img_w, img_h, theme)
    else:
        _image_placeholder(prs_slide, img_left, _MARGIN, img_w, img_h, theme)

    if slide.notes:
        prs_slide.notes_slide.notes_text_frame.text = slide.notes


# ---------------------------------------------------------------------------
# MinimalDeck layout dispatch (compat path)
# ---------------------------------------------------------------------------

_MINIMAL_LAYOUT_FNS = {
    "title": _layout_title_slide,
    "bullets": _layout_bullets_slide,
    "section": _layout_section_slide,
    "image_right": _layout_image_right_slide,
}
