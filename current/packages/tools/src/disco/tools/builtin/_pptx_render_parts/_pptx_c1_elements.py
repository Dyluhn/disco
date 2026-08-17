"""C1 Element renderer — maps each Element to a python-pptx shape.

Moved verbatim out of ``_pptx_render.py``.
"""

from __future__ import annotations

from disco.core.brand.tokens import Theme

from .._deck_schema import Element, Slide
from ._geometry import _MARGIN, _SLIDE_H
from ._pptx_primitives import (
    _add_accent_bar,
    _add_textbox,
    _image_placeholder,
    _rgb_from_hex,
    _set_run_style,
)


def _is_bullet_el(el: Element) -> bool:
    """A bullet body Element: a text Element whose text carries the "• " prefix.

    The C1 lowerer (`_deck_schema._layout_bullets` / `_layout_image_right` /
    `_layout_two_column` / `_layout_comparison`) emits each bullet as an
    independent absolutely-positioned textbox with this prefix; W-23 regroups
    them at render time so a wrapped 2nd line cannot spill into the next box.
    """
    return el.kind == "text" and el.text.startswith("• ")


def _render_bullet_group(prs_slide, group: list[Element], theme: Theme) -> None:  # type: ignore[type-arg]
    """W-23: render consecutive bullet Elements as ONE auto-fit text frame —
    one PARAGRAPH per bullet — instead of one absolute textbox per bullet.

    The per-bullet model (`top = body_top + i*line_h`, `height = line_h*1.15`)
    reserves NO space for a wrapped 2nd line; under font substitution in
    LibreOffice / Google Slides that 2nd line spilled into the next bullet's
    absolute box → overlap. Grouping into one flowing frame (word_wrap + TOP
    anchor + zero margins + small `space_after`, autofit via `_add_textbox`)
    lets bullets flow and never overlap. Text keeps its "• " prefix so the
    deck stays editable as plain paragraphs.
    """
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Pt

    _ALIGN_MAP = {"LEFT": PP_ALIGN.LEFT, "CENTER": PP_ALIGN.CENTER, "RIGHT": PP_ALIGN.RIGHT}

    first = group[0]
    left = int(first.left)
    width = int(first.width)
    top = int(first.top)
    # Span down to the bottom safe-area margin so wrapped bullets have room to
    # flow; normAutofit (set in _add_textbox) only shrinks if even that overflows.
    height = max(int(first.height), (_SLIDE_H - _MARGIN) - top)

    tf = _add_textbox(prs_slide, left, top, width, height)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)

    for i, el in enumerate(group):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN_MAP.get(el.align, PP_ALIGN.LEFT)
        p.line_spacing = 1.05
        p.space_after = Pt(6)
        run = p.add_run()
        _set_run_style(
            run,
            el.text,
            el.font_name or "Helvetica",
            el.font_size_pt,
            el.hex_color or "#1a1813",
            bold=el.bold,
            italic=el.italic,
        )


def _render_slide_elements(prs_slide, slide: Slide, theme: Theme) -> None:  # type: ignore[type-arg]
    """Render a slide's Elements, grouping consecutive bullet text Elements that
    share the same left/width into ONE auto-fit frame (W-23).

    Same-left/width grouping keeps the left vs right columns of two_column /
    comparison layouts as separate frames (their bullets differ by `left`), and
    leaves non-bullet elements (title, accent bar, image, labels) untouched.
    """
    group: list[Element] = []

    def _flush() -> None:
        if group:
            _render_bullet_group(prs_slide, list(group), theme)
            group.clear()

    for el in slide.elements:
        if _is_bullet_el(el) and (
            not group
            or (int(el.left) == int(group[0].left) and int(el.width) == int(group[0].width))
        ):
            group.append(el)
            continue
        _flush()
        if _is_bullet_el(el):
            group.append(el)
        else:
            _render_element(prs_slide, el, theme)
    _flush()


def _render_element(prs_slide, el: Element, theme: Theme) -> None:  # type: ignore[type-arg]
    """Render one C1 Element to a python-pptx slide shape."""

    if el.kind == "text":
        _render_text_element(prs_slide, el)
    elif el.kind == "accent_bar":
        _render_accent_bar_element(prs_slide, el)
    elif el.kind == "image":
        _render_image_element(prs_slide, el, theme)
    elif el.kind == "rect":
        _render_rect_element(prs_slide, el)
    # unknown kinds are skipped cleanly (no crash)


def _render_text_element(prs_slide, el: Element) -> None:  # type: ignore[type-arg]
    """Render a text Element as an absolutely-positioned textbox."""
    from pptx.enum.text import PP_ALIGN

    _ALIGN_MAP = {"LEFT": PP_ALIGN.LEFT, "CENTER": PP_ALIGN.CENTER, "RIGHT": PP_ALIGN.RIGHT}
    tf = _add_textbox(
        prs_slide,
        int(el.left),
        int(el.top),
        int(el.width),
        int(el.height),
    )
    tf.word_wrap = el.word_wrap
    p = tf.paragraphs[0]
    p.alignment = _ALIGN_MAP.get(el.align, PP_ALIGN.LEFT)
    run = p.add_run()
    # NOTE: python-pptx only stores the font *name* in the slide XML (e.g.
    # "Fraunces"). PPTX viewers (PowerPoint, LibreOffice) render it only if the
    # font is installed on the host OS. The base64 data-URI embedding in
    # font_face_css() fixes HTML/PDF output; PPTX fonts are NOT fixed here.
    _set_run_style(
        run,
        el.text,
        el.font_name or "Helvetica",
        el.font_size_pt,
        el.hex_color or "#1a1813",
        bold=el.bold,
        italic=el.italic,
    )


def _render_accent_bar_element(prs_slide, el: Element) -> None:  # type: ignore[type-arg]
    """Render an accent_bar Element as a thin filled rectangle."""
    _add_accent_bar(
        prs_slide,
        int(el.left),
        int(el.top),
        int(el.width),
        el.fill_hex,
    )


def _render_image_element(prs_slide, el: Element, theme: Theme) -> None:  # type: ignore[type-arg]
    """Render an image Element; falls back to styled placeholder when there is
    neither generated bytes nor a readable path."""
    import io

    from pptx.util import Emu

    # C7 wire: prefer the generated bytes (sandbox-agnostic) over a path on disk.
    source: object | None = io.BytesIO(el.image_bytes) if el.image_bytes else el.image_path
    if source is not None:
        try:
            prs_slide.shapes.add_picture(
                source,
                Emu(int(el.left)),
                Emu(int(el.top)),
                Emu(int(el.width)),
                Emu(int(el.height)),
            )
            return
        except Exception:
            pass  # fall through to placeholder

    # Styled placeholder box (visible "image" label)
    _image_placeholder(
        prs_slide,
        int(el.left),
        int(el.top),
        int(el.width),
        int(el.height),
        theme,
    )


def _render_rect_element(prs_slide, el: Element) -> None:  # type: ignore[type-arg]
    """Render a rect Element as a filled rectangle shape."""
    from pptx.util import Emu

    box = prs_slide.shapes.add_shape(
        1,
        Emu(int(el.left)),
        Emu(int(el.top)),
        Emu(int(el.width)),
        Emu(int(el.height)),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb_from_hex(el.fill_hex)
    if el.border_hex:
        box.line.color.rgb = _rgb_from_hex(el.border_hex)
    else:
        box.line.fill.background()
